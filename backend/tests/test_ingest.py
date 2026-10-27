import io
from datetime import UTC, datetime, timedelta
from pathlib import Path
from uuid import UUID, uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from app.config import Settings
from app.main import app
from app.models.course import Course, CourseMember, CourseMemberRole
from app.models.document import Chunk, Document, IngestJob, IngestJobStatus
from app.models.user import User, UserRole
from app.services.chunk import CHUNKING_VERSION
from app.services.embedder import FakeEmbedder, get_embedder
from app.services.ingest import (
    COVERAGE_INCOMPLETE,
    LOW_EXTRACTED_TEXT,
    IngestLeaseLost,
    _touch_running_job,
    process_ingest_job,
    retry_ingest_job,
)
from app.services.parse import ParseLimits, parse_document
from app.services.upstream import UPSTREAM_UNAVAILABLE, UpstreamUnavailable
from fastapi.testclient import TestClient
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from tests.conftest import add_user, enable_retrieval
from tests.test_parse import make_pdf

UNIQUE_SENTENCE = "紫石英口令：本课向量必须绑定 course_id-alpha-9371。"
PPTX_SENTENCE = "紫石英PPT口令：幻灯片锚点 delta-8820。"
DOCX_SENTENCE = "紫石英Word口令：实验指导锚点 epsilon-5503。"
SAMPLE_MD = Path(__file__).parent / "fixtures" / "sample.md"
SAMPLE_PPTX = Path(__file__).parent / "fixtures" / "sample.pptx"
SAMPLE_DOCX = Path(__file__).parent / "fixtures" / "sample.docx"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class RecordingEmbedder(FakeEmbedder):
    def __init__(self, *, fail_on_call: int | None = None) -> None:
        super().__init__()
        self.batch_sizes: list[int] = []
        self.fail_on_call = fail_on_call

    def embed(self, texts: list[str]) -> list[list[float]]:
        self.batch_sizes.append(len(texts))
        if self.fail_on_call == len(self.batch_sizes):
            raise RuntimeError("scripted embedding failure")
        return super().embed(texts)


def login_as(client: TestClient, username: str, password: str) -> None:
    client.post("/auth/logout")
    response = client.post("/auth/login", json={"username": username, "password": password})
    assert response.status_code == 200


def add_course(db: Session, title: str) -> Course:
    course = Course(title=title)
    db.add(course)
    db.commit()
    db.refresh(course)
    return course


def add_member(db: Session, course: Course, user: User, role: CourseMemberRole) -> None:
    db.add(CourseMember(course_id=course.id, user_id=user.id, role=role.value))
    db.commit()


def make_zip(entries: list[tuple[str, bytes]]) -> bytes:
    output = io.BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        for name, content in entries:
            archive.writestr(name, content)
    return output.getvalue()


def upload_as_teacher(client: TestClient, db: Session, filename: str, content: bytes) -> UUID:
    teacher = add_user(db, "t1", "pass", UserRole.teacher)
    course = add_course(db, "Course 1")
    add_member(db, course, teacher, CourseMemberRole.teacher)
    login_as(client, "t1", "pass")
    uploaded = client.post(
        f"/courses/{course.id}/documents",
        files={"file": (filename, content, "application/octet-stream")},
    )
    assert uploaded.status_code == 201
    return UUID(uploaded.json()["id"])


def assert_safe_quality_metadata(payload: dict, chunks: list[Chunk]) -> None:
    assert payload["source_unit_count"] >= 1
    assert payload["extracted_char_count"] >= 1
    assert payload["chunk_count"] == len(chunks)
    assert payload["quality_warning"] in {None, LOW_EXTRACTED_TEXT, COVERAGE_INCOMPLETE}
    assert payload["chunking_version"] == CHUNKING_VERSION
    assert all(chunk.source_metadata is not None for chunk in chunks)
    assert {"text", "snippet", "storage_path"}.isdisjoint(payload)



def test_upload_rejects_over_100mb(client: TestClient, db: Session) -> None:
    teacher = add_user(db, "t1", "pass", UserRole.teacher)
    course = add_course(db, "Course 1")
    add_member(db, course, teacher, CourseMemberRole.teacher)
    login_as(client, "t1", "pass")
    # One byte over the 100 MiB gate.
    oversized = b"x" * (100 * 1024 * 1024 + 1)
    uploaded = client.post(
        f"/courses/{course.id}/documents",
        files={"file": ("huge.pdf", oversized, "application/pdf")},
    )
    assert uploaded.status_code == 413
    assert uploaded.json()["detail"] == "file too large"


def test_upload_limit_constant_is_100mb() -> None:
    from app.api.documents import _MAX_UPLOAD_BYTES

    assert _MAX_UPLOAD_BYTES == 100 * 1024 * 1024


def test_ooxml_uncompressed_default_is_at_least_100mb() -> None:
    settings = Settings(_env_file=None)
    assert settings.ooxml_max_uncompressed_bytes >= 100 * 1024 * 1024
    assert settings.ooxml_max_entry_bytes >= 50 * 1024 * 1024


@pytest.mark.parametrize("threshold", [0, 5_001])
def test_low_text_threshold_is_bounded(threshold: int) -> None:
    with pytest.raises(ValidationError):
        Settings(_env_file=None, ingest_low_text_chars=threshold)


def test_upload_sample_md_indexes_unique_sentence(client: TestClient, db: Session) -> None:
    teacher = add_user(db, "t1", "pass", UserRole.teacher)
    course = add_course(db, "Course 1")
    add_member(db, course, teacher, CourseMemberRole.teacher)
    login_as(client, "t1", "pass")

    content = SAMPLE_MD.read_bytes()
    assert UNIQUE_SENTENCE.encode("utf-8") in content
    uploaded = client.post(
        f"/courses/{course.id}/documents",
        files={"file": ("sample.md", content, "text/markdown")},
    )
    assert uploaded.status_code == 201
    doc_id = UUID(uploaded.json()["id"])

    detail = client.get(f"/courses/{course.id}/documents/{doc_id}")
    assert detail.status_code == 200
    assert detail.json()["job_status"] == "succeeded"
    assert detail.json()["error_message"] is None

    db.expire_all()
    chunks = db.execute(select(Chunk).where(Chunk.document_id == doc_id)).scalars().all()
    assert chunks
    assert all(chunk.course_id == course.id for chunk in chunks)
    assert any(UNIQUE_SENTENCE in chunk.text for chunk in chunks)
    assert all(len(chunk.embedding) == 1024 for chunk in chunks)
    assert all(chunk.embedding_model == FakeEmbedder().index_identity for chunk in chunks)
    assert_safe_quality_metadata(detail.json(), chunks)


def test_upload_sample_pptx_indexes_unique_sentence(client: TestClient, db: Session) -> None:
    teacher = add_user(db, "t1", "pass", UserRole.teacher)
    course = add_course(db, "Course 1")
    add_member(db, course, teacher, CourseMemberRole.teacher)
    login_as(client, "t1", "pass")

    content = SAMPLE_PPTX.read_bytes()
    uploaded = client.post(
        f"/courses/{course.id}/documents",
        files={"file": ("sample.pptx", content, PPTX_MIME)},
    )
    assert uploaded.status_code == 201
    doc_id = UUID(uploaded.json()["id"])

    detail = client.get(f"/courses/{course.id}/documents/{doc_id}")
    assert detail.status_code == 200
    assert detail.json()["job_status"] == "succeeded"
    assert detail.json()["error_message"] is None

    db.expire_all()
    chunks = db.execute(select(Chunk).where(Chunk.document_id == doc_id)).scalars().all()
    assert chunks
    assert all(chunk.course_id == course.id for chunk in chunks)
    assert any(PPTX_SENTENCE in chunk.text for chunk in chunks)
    assert all(len(chunk.embedding) == 1024 for chunk in chunks)
    assert_safe_quality_metadata(detail.json(), chunks)


def test_corrupt_pptx_fails_without_chunks(client: TestClient, db: Session) -> None:
    teacher = add_user(db, "t1", "pass", UserRole.teacher)
    course = add_course(db, "Course 1")
    add_member(db, course, teacher, CourseMemberRole.teacher)
    login_as(client, "t1", "pass")

    uploaded = client.post(
        f"/courses/{course.id}/documents",
        files={"file": ("broken.pptx", b"PK\x03\x04-not-a-real-pptx", PPTX_MIME)},
    )
    assert uploaded.status_code == 201
    doc_id = UUID(uploaded.json()["id"])
    detail = client.get(f"/courses/{course.id}/documents/{doc_id}")
    assert detail.status_code == 200
    assert detail.json()["job_status"] == "failed"
    assert detail.json()["error_message"]

    db.expire_all()
    chunks = db.execute(select(Chunk).where(Chunk.document_id == doc_id)).scalars().all()
    assert chunks == []


def test_upload_sample_docx_indexes_unique_sentence(client: TestClient, db: Session) -> None:
    teacher = add_user(db, "t1", "pass", UserRole.teacher)
    course = add_course(db, "Course 1")
    add_member(db, course, teacher, CourseMemberRole.teacher)
    login_as(client, "t1", "pass")

    content = SAMPLE_DOCX.read_bytes()
    uploaded = client.post(
        f"/courses/{course.id}/documents",
        files={"file": ("sample.docx", content, DOCX_MIME)},
    )
    assert uploaded.status_code == 201
    doc_id = UUID(uploaded.json()["id"])

    detail = client.get(f"/courses/{course.id}/documents/{doc_id}")
    assert detail.status_code == 200
    assert detail.json()["job_status"] == "succeeded"
    assert detail.json()["error_message"] is None

    db.expire_all()
    chunks = db.execute(select(Chunk).where(Chunk.document_id == doc_id)).scalars().all()
    assert chunks
    assert all(chunk.course_id == course.id for chunk in chunks)
    assert any(DOCX_SENTENCE in chunk.text for chunk in chunks)
    assert all(len(chunk.embedding) == 1024 for chunk in chunks)
    assert_safe_quality_metadata(detail.json(), chunks)


def test_upload_sample_pdf_stays_available(client: TestClient, db: Session) -> None:
    teacher = add_user(db, "t1", "pass", UserRole.teacher)
    course = add_course(db, "Course 1")
    add_member(db, course, teacher, CourseMemberRole.teacher)
    login_as(client, "t1", "pass")

    uploaded = client.post(
        f"/courses/{course.id}/documents",
        files={"file": ("sample.pdf", make_pdf("Sample PDF"), "application/pdf")},
    )

    assert uploaded.status_code == 201
    doc_id = UUID(uploaded.json()["id"])
    detail = client.get(f"/courses/{course.id}/documents/{doc_id}")
    assert detail.json()["job_status"] == IngestJobStatus.succeeded.value
    chunks = db.execute(select(Chunk).where(Chunk.document_id == doc_id)).scalars().all()
    assert chunks
    assert_safe_quality_metadata(detail.json(), chunks)


def test_short_text_succeeds_with_quality_warning(
    client: TestClient, db: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("INGEST_LOW_TEXT_CHARS", "100")
    doc_id = upload_as_teacher(client, db, "brief.txt", b"brief")
    document = db.get(Document, doc_id)

    detail = client.get(f"/courses/{document.course_id}/documents/{doc_id}")

    assert detail.status_code == 200
    assert detail.json() == {
        "id": str(doc_id),
        "filename": "brief.txt",
        "folder_id": None,
        "folder_name": None,
        "retrieval_enabled": False,
        "job_status": IngestJobStatus.succeeded.value,
        "error_message": None,
        "source_unit_count": 1,
        "extracted_char_count": 5,
        "chunk_count": 1,
        "quality_warning": LOW_EXTRACTED_TEXT,
        "chunking_version": CHUNKING_VERSION,
    }


def test_empty_extracted_text_fails_without_quality_metadata(
    client: TestClient, db: Session
) -> None:
    doc_id = upload_as_teacher(client, db, "blank.txt", b"   \n")
    document = db.get(Document, doc_id)

    detail = client.get(f"/courses/{document.course_id}/documents/{doc_id}")

    assert detail.status_code == 200
    assert detail.json()["job_status"] == IngestJobStatus.failed.value
    assert detail.json()["error_message"] == "no extractable text"
    assert detail.json()["source_unit_count"] is None
    assert detail.json()["extracted_char_count"] is None
    assert detail.json()["chunk_count"] is None
    assert detail.json()["quality_warning"] is None
    assert db.execute(select(Chunk).where(Chunk.document_id == doc_id)).scalars().all() == []


@pytest.mark.parametrize(
    ("filename", "environment", "content"),
    [
        (
            "too-many.pptx",
            {"OOXML_MAX_ZIP_ENTRIES": "10"},
            make_zip([(f"entry-{index}.xml", b"x") for index in range(11)]),
        ),
        (
            "too-large-total.docx",
            {"OOXML_MAX_UNCOMPRESSED_BYTES": str(1024 * 1024)},
            make_zip([("one.xml", b"a" * (600 * 1024)), ("two.xml", b"b" * (600 * 1024))]),
        ),
        (
            "too-large-entry.pptx",
            {"OOXML_MAX_ENTRY_BYTES": str(1024 * 1024)},
            make_zip([("large.xml", b"a" * (1024 * 1024 + 1))]),
        ),
        (
            "too-compressible.docx",
            {"OOXML_MAX_COMPRESSION_RATIO": "2"},
            make_zip([("compressed.xml", b"a" * (1024 * 1024))]),
        ),
    ],
    ids=["entry-count", "total-uncompressed", "entry-size", "compression-ratio"],
)
def test_ooxml_zip_resource_limits_fail_without_chunks(
    client: TestClient,
    db: Session,
    monkeypatch: pytest.MonkeyPatch,
    filename: str,
    environment: dict[str, str],
    content: bytes,
) -> None:
    for key, value in environment.items():
        monkeypatch.setenv(key, value)
    doc_id = upload_as_teacher(client, db, filename, content)

    detail = client.get(f"/courses/{db.get(Document, doc_id).course_id}/documents/{doc_id}")

    assert detail.status_code == 200
    assert detail.json()["job_status"] == IngestJobStatus.failed.value
    assert detail.json()["error_message"] == "OOXML resource limit exceeded"
    assert db.execute(select(Chunk).where(Chunk.document_id == doc_id)).scalars().all() == []


def test_ooxml_resource_limit_preserves_previous_complete_chunks(
    db: Session, tmp_path: Path
) -> None:
    course = add_course(db, "Course 1")
    source = tmp_path / "too-many.pptx"
    source.write_bytes(make_zip([(f"entry-{index}.xml", b"x") for index in range(11)]))
    document = Document(
        course_id=course.id,
        uploaded_by=uuid4(),
        filename=source.name,
        storage_path=str(source),
        source_unit_count=2,
        extracted_char_count=800,
        chunk_count=1,
        quality_warning=None,
    )
    embedder = FakeEmbedder()
    db.add(document)
    db.flush()
    old_chunk = Chunk(
        course_id=course.id,
        document_id=document.id,
        page=1,
        text="previous complete index",
        embedding=embedder.embed(["previous complete index"])[0],
        embedding_model=embedder.index_identity,
    )
    job = IngestJob(document_id=document.id, status=IngestJobStatus.pending.value)
    db.add_all([old_chunk, job])
    db.commit()

    process_ingest_job(
        db,
        job.id,
        embedder,
        parse_limits=ParseLimits(max_zip_entries=10),
    )

    db.expire_all()
    stored_job = db.get(IngestJob, job.id)
    chunks = db.execute(select(Chunk).where(Chunk.document_id == document.id)).scalars().all()
    assert stored_job is not None
    assert stored_job.status == IngestJobStatus.failed.value
    assert stored_job.error_message == "OOXML resource limit exceeded"
    assert [chunk.id for chunk in chunks] == [old_chunk.id]
    stored_document = db.get(Document, document.id)
    assert stored_document.source_unit_count == 2
    assert stored_document.extracted_char_count == 800
    assert stored_document.chunk_count == 1
    assert stored_document.quality_warning is None


@pytest.mark.parametrize("filename", ["large.pptx", "large.docx"])
def test_ooxml_extracted_text_limit_fails_without_chunks(
    client: TestClient,
    db: Session,
    monkeypatch: pytest.MonkeyPatch,
    filename: str,
) -> None:
    monkeypatch.setenv("PARSE_MAX_TEXT_CHARS", "10000")
    text = "x" * 10_001
    output = io.BytesIO()
    if filename.endswith(".pptx"):
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.add_textbox(0, 0, 1, 1).text_frame.text = text
        presentation.save(output)
    else:
        from docx import Document as WordDocument

        document = WordDocument()
        document.add_paragraph(text)
        document.save(output)
    doc_id = upload_as_teacher(client, db, filename, output.getvalue())

    detail = client.get(f"/courses/{db.get(Document, doc_id).course_id}/documents/{doc_id}")

    assert detail.status_code == 200
    assert detail.json()["job_status"] == IngestJobStatus.failed.value
    assert detail.json()["error_message"] == "extracted text exceeds resource limit"
    assert db.execute(select(Chunk).where(Chunk.document_id == doc_id)).scalars().all() == []


def test_pptx_slide_limit_fails_without_chunks(
    client: TestClient, db: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("PARSE_MAX_PAGES", "1")
    output = io.BytesIO()
    from pptx import Presentation

    presentation = Presentation()
    for _ in range(2):
        presentation.slides.add_slide(presentation.slide_layouts[5])
    presentation.save(output)
    doc_id = upload_as_teacher(client, db, "two-slides.pptx", output.getvalue())

    detail = client.get(f"/courses/{db.get(Document, doc_id).course_id}/documents/{doc_id}")

    assert detail.status_code == 200
    assert detail.json()["job_status"] == IngestJobStatus.failed.value
    assert detail.json()["error_message"] == "extracted text exceeds resource limit"
    assert db.execute(select(Chunk).where(Chunk.document_id == doc_id)).scalars().all() == []


def test_corrupt_docx_fails_without_chunks(client: TestClient, db: Session) -> None:
    teacher = add_user(db, "t1", "pass", UserRole.teacher)
    course = add_course(db, "Course 1")
    add_member(db, course, teacher, CourseMemberRole.teacher)
    login_as(client, "t1", "pass")

    uploaded = client.post(
        f"/courses/{course.id}/documents",
        files={"file": ("broken.docx", b"PK\x03\x04-not-a-real-docx", DOCX_MIME)},
    )
    assert uploaded.status_code == 201
    doc_id = UUID(uploaded.json()["id"])
    detail = client.get(f"/courses/{course.id}/documents/{doc_id}")
    assert detail.status_code == 200
    assert detail.json()["job_status"] == "failed"
    assert detail.json()["error_message"]

    db.expire_all()
    chunks = db.execute(select(Chunk).where(Chunk.document_id == doc_id)).scalars().all()
    assert chunks == []


def test_corrupt_pdf_fails_without_chunks(client: TestClient, db: Session) -> None:
    teacher = add_user(db, "t1", "pass", UserRole.teacher)
    course = add_course(db, "Course 1")
    add_member(db, course, teacher, CourseMemberRole.teacher)
    login_as(client, "t1", "pass")

    uploaded = client.post(
        f"/courses/{course.id}/documents",
        files={"file": ("broken.pdf", b"%PDF-not-a-real-file", "application/pdf")},
    )
    assert uploaded.status_code == 201
    doc_id = UUID(uploaded.json()["id"])
    detail = client.get(f"/courses/{course.id}/documents/{doc_id}")
    assert detail.status_code == 200
    assert detail.json()["job_status"] == "failed"
    assert detail.json()["error_message"]

    db.expire_all()
    chunks = db.execute(select(Chunk).where(Chunk.document_id == doc_id)).scalars().all()
    assert chunks == []


@pytest.mark.parametrize(
    ("filename", "content_type"),
    [
        ("legacy.ppt", "application/vnd.ms-powerpoint"),
        ("legacy.doc", "application/msword"),
    ],
)
def test_legacy_binary_office_is_rejected_before_persistence(
    client: TestClient,
    db: Session,
    filename: str,
    content_type: str,
) -> None:
    teacher = add_user(db, "t1", "pass", UserRole.teacher)
    course = add_course(db, "Course 1")
    add_member(db, course, teacher, CourseMemberRole.teacher)
    login_as(client, "t1", "pass")

    uploaded = client.post(
        f"/courses/{course.id}/documents",
        files={"file": (filename, b"legacy binary content", content_type)},
    )

    assert uploaded.status_code == 400
    assert uploaded.json() == {
        "detail": "unsupported file type; accepted: PDF, PPTX, DOCX, Markdown, or plain text"
    }
    assert db.execute(select(Document).where(Document.course_id == course.id)).scalars().all() == []
    assert db.execute(select(IngestJob)).scalars().all() == []


def test_student_cannot_delete_document(client: TestClient, db: Session) -> None:
    student = add_user(db, "alice", "pass", UserRole.student)
    teacher = add_user(db, "t1", "pass", UserRole.teacher)
    course = add_course(db, "Course 1")
    add_member(db, course, teacher, CourseMemberRole.teacher)
    add_member(db, course, student, CourseMemberRole.student)
    login_as(client, "t1", "pass")
    uploaded = client.post(
        f"/courses/{course.id}/documents",
        files={"file": ("sample.md", SAMPLE_MD.read_bytes(), "text/markdown")},
    )
    assert uploaded.status_code == 201
    doc_id = UUID(uploaded.json()["id"])
    login_as(client, "alice", "pass")
    deleted = client.delete(f"/courses/{course.id}/documents/{doc_id}")
    assert deleted.status_code == 403


def test_student_cannot_upload(client: TestClient, db: Session) -> None:
    student = add_user(db, "alice", "pass", UserRole.student)
    teacher = add_user(db, "t1", "pass", UserRole.teacher)
    course = add_course(db, "Course 1")
    add_member(db, course, teacher, CourseMemberRole.teacher)
    add_member(db, course, student, CourseMemberRole.student)
    login_as(client, "alice", "pass")
    uploaded = client.post(
        f"/courses/{course.id}/documents",
        files={"file": ("sample.md", SAMPLE_MD.read_bytes(), "text/markdown")},
    )
    assert uploaded.status_code == 403


def test_student_member_can_read_documents_and_nonmember_is_denied(
    client: TestClient, db: Session
) -> None:
    student = add_user(db, "alice", "pass", UserRole.student)
    add_user(db, "mallory", "pass", UserRole.student)
    teacher = add_user(db, "t1", "pass", UserRole.teacher)
    course = add_course(db, "Course 1")
    other_course = add_course(db, "Course 2")
    add_member(db, course, teacher, CourseMemberRole.teacher)
    add_member(db, course, student, CourseMemberRole.student)
    add_member(db, other_course, student, CourseMemberRole.student)
    login_as(client, "t1", "pass")
    content = SAMPLE_MD.read_bytes()
    uploaded = client.post(
        f"/courses/{course.id}/documents",
        files={"file": ("sample.md", content, "text/markdown")},
    )
    assert uploaded.status_code == 201
    doc_id = uploaded.json()["id"]

    login_as(client, "alice", "pass")
    assert client.get(f"/courses/{course.id}/documents").json() == []
    assert client.get(f"/courses/{course.id}/documents/{doc_id}").status_code == 404
    assert client.get(f"/courses/{course.id}/documents/{doc_id}/content").status_code == 404
    enable_retrieval(db, doc_id)
    listed = client.get(f"/courses/{course.id}/documents")
    assert listed.status_code == 200
    assert [row["id"] for row in listed.json()] == [doc_id]
    detail = client.get(f"/courses/{course.id}/documents/{doc_id}")
    assert detail.status_code == 200
    assert detail.json()["filename"] == "sample.md"
    served = client.get(f"/courses/{course.id}/documents/{doc_id}/content")
    assert served.status_code == 200
    assert content in served.content

    # Same document id under another course the student belongs to must not leak.
    cross_detail = client.get(f"/courses/{other_course.id}/documents/{doc_id}")
    assert cross_detail.status_code == 404
    cross_content = client.get(f"/courses/{other_course.id}/documents/{doc_id}/content")
    assert cross_content.status_code == 404
    assert client.get(f"/courses/{other_course.id}/documents").json() == []

    login_as(client, "mallory", "pass")
    assert client.get(f"/courses/{course.id}/documents").status_code == 403
    assert client.get(f"/courses/{course.id}/documents/{doc_id}").status_code == 403
    assert client.get(f"/courses/{course.id}/documents/{doc_id}/content").status_code == 403


def test_ingest_sends_configured_embedding_batches(
    client: TestClient, db: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    teacher = add_user(db, "t1", "pass", UserRole.teacher)
    course = add_course(db, "Course 1")
    add_member(db, course, teacher, CourseMemberRole.teacher)
    login_as(client, "t1", "pass")
    embedder = RecordingEmbedder()
    app.dependency_overrides[get_embedder] = lambda: embedder
    monkeypatch.setenv("EMBEDDING_BATCH_SIZE", "2")

    uploaded = client.post(
        f"/courses/{course.id}/documents",
        files={"file": ("large.txt", ("batch boundary " * 200).encode(), "text/plain")},
    )

    assert uploaded.status_code == 201
    assert len(embedder.batch_sizes) > 1
    assert all(size <= 2 for size in embedder.batch_sizes)
    assert sum(embedder.batch_sizes) > 2


def test_mid_batch_failure_preserves_previous_complete_chunks(
    db: Session, tmp_path: Path
) -> None:
    course = add_course(db, "Course 1")
    source = tmp_path / "retry.txt"
    source.write_text("transactional chunks " * 150, encoding="utf-8")
    document = Document(
        course_id=course.id,
        uploaded_by=uuid4(),
        filename=source.name,
        storage_path=str(source),
        source_unit_count=2,
        extracted_char_count=800,
        chunk_count=1,
        quality_warning=None,
    )
    db.add(document)
    db.flush()
    embedder = RecordingEmbedder(fail_on_call=2)
    old_chunk = Chunk(
        course_id=course.id,
        document_id=document.id,
        page=1,
        text="previous complete index",
        embedding=embedder.embed(["previous complete index"])[0],
        embedding_model=embedder.index_identity,
    )
    embedder.batch_sizes.clear()
    job = IngestJob(document_id=document.id, status=IngestJobStatus.pending.value)
    db.add_all([old_chunk, job])
    db.commit()

    process_ingest_job(db, job.id, embedder, batch_size=1)

    db.expire_all()
    stored_job = db.get(IngestJob, job.id)
    chunks = db.execute(select(Chunk).where(Chunk.document_id == document.id)).scalars().all()
    assert stored_job is not None
    assert stored_job.status == IngestJobStatus.failed.value
    assert stored_job.run_token is None
    assert [chunk.id for chunk in chunks] == [old_chunk.id]
    assert chunks[0].text == "previous complete index"
    stored_document = db.get(Document, document.id)
    assert stored_document.source_unit_count == 2
    assert stored_document.extracted_char_count == 800
    assert stored_document.chunk_count == 1
    assert stored_document.quality_warning is None


def test_ingest_strips_postgresql_nul_before_embedding_and_storage(
    db: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    course = add_course(db, "Course 1")
    source = tmp_path / "nul-source.pdf"
    source.write_bytes(b"placeholder")
    document = Document(
        course_id=course.id,
        uploaded_by=uuid4(),
        filename=source.name,
        storage_path=str(source),
    )
    db.add(document)
    db.flush()
    job = IngestJob(document_id=document.id, status=IngestJobStatus.pending.value)
    db.add(job)
    db.commit()
    from app.services.parse_quality import build_native_pdf_quality_report
    text = "vector\x00database evidence"
    report = build_native_pdf_quality_report(
        [(1, text), (2, "\x00")], parser_version="test", input_sha256="0" * 64,
    )
    monkeypatch.setattr(
        "app.services.ingest.parse_document_with_quality",
        lambda *args, **kwargs: ([(1, text), (2, "\x00")], report),
    )

    class CapturingEmbedder(FakeEmbedder):
        def embed(self, texts):
            assert texts == ["vectordatabase evidence"]
            return super().embed(texts)

    process_ingest_job(db, job.id, CapturingEmbedder())

    db.expire_all()
    stored_job = db.get(IngestJob, job.id)
    chunks = db.execute(select(Chunk).where(Chunk.document_id == document.id)).scalars().all()
    stored_document = db.get(Document, document.id)
    assert stored_job is not None
    assert stored_job.status == IngestJobStatus.succeeded.value
    assert [chunk.text for chunk in chunks] == ["vectordatabase evidence"]
    assert stored_document is not None
    assert stored_document.source_unit_count == 1
    assert stored_document.quality_warning == COVERAGE_INCOMPLETE
    assert stored_document.extracted_char_count == len("vectordatabase evidence")


def test_upstream_embedding_failure_is_safe_and_explicitly_retryable(
    db: Session, tmp_path: Path
) -> None:
    course = add_course(db, "Course 1")
    source = tmp_path / "retry.txt"
    source.write_text("safe retry content " * 150, encoding="utf-8")
    document = Document(
        course_id=course.id,
        uploaded_by=uuid4(),
        filename=source.name,
        storage_path=str(source),
    )
    db.add(document)
    db.flush()
    working_embedder = FakeEmbedder()
    old_chunk = Chunk(
        course_id=course.id,
        document_id=document.id,
        page=1,
        text="previous complete index",
        embedding=working_embedder.embed(["previous complete index"])[0],
        embedding_model=working_embedder.index_identity,
    )
    job = IngestJob(document_id=document.id, status=IngestJobStatus.pending.value)
    db.add_all([old_chunk, job])
    db.commit()

    class FailingEmbedder(FakeEmbedder):
        def embed(self, texts: list[str]) -> list[list[float]]:
            raise UpstreamUnavailable("provider-body-should-not-be-stored")

    process_ingest_job(db, job.id, FailingEmbedder())

    db.expire_all()
    stored_job = db.get(IngestJob, job.id)
    chunks = db.execute(select(Chunk).where(Chunk.document_id == document.id)).scalars().all()
    assert stored_job is not None
    assert stored_job.status == IngestJobStatus.failed.value
    assert stored_job.error_message == UPSTREAM_UNAVAILABLE
    assert [chunk.id for chunk in chunks] == [old_chunk.id]
    assert retry_ingest_job(db, job.id, datetime.now(UTC) - timedelta(minutes=15))
    db.expire_all()
    assert db.get(IngestJob, job.id).status == IngestJobStatus.pending.value


@pytest.mark.parametrize(
    "recoverable_status",
    [
        IngestJobStatus.failed.value,
        IngestJobStatus.pending.value,
        IngestJobStatus.running.value,
    ],
)
def test_teacher_can_recover_failed_or_stale_ingest_job(
    client: TestClient,
    db: Session,
    recoverable_status: str,
) -> None:
    teacher = add_user(db, f"teacher-{recoverable_status}", "pass", UserRole.teacher)
    course = add_course(db, "Course 1")
    add_member(db, course, teacher, CourseMemberRole.teacher)
    document = Document(
        course_id=course.id,
        uploaded_by=teacher.id,
        filename=SAMPLE_MD.name,
        storage_path=str(SAMPLE_MD),
    )
    db.add(document)
    db.flush()
    job = IngestJob(
        document_id=document.id,
        status=recoverable_status,
        run_token=(
            uuid4() if recoverable_status == IngestJobStatus.running.value else None
        ),
        updated_at=(
            datetime.now(UTC)
            if recoverable_status == IngestJobStatus.failed.value
            else datetime.now(UTC) - timedelta(hours=1)
        ),
    )
    db.add(job)
    db.commit()
    login_as(client, teacher.username, "pass")

    retried = client.post(f"/courses/{course.id}/documents/{document.id}/retry")

    assert retried.status_code == 202
    db.expire_all()
    stored_job = db.get(IngestJob, job.id)
    chunks = db.execute(select(Chunk).where(Chunk.document_id == document.id)).scalars().all()
    assert stored_job is not None
    assert stored_job.status == IngestJobStatus.succeeded.value
    assert stored_job.run_token is None
    assert chunks
    assert db.get(Document, document.id).chunk_count == len(chunks)
    chunk_ids = [chunk.id for chunk in chunks]

    duplicate = client.post(f"/courses/{course.id}/documents/{document.id}/retry")
    assert duplicate.status_code == 409
    db.expire_all()
    after_duplicate = db.execute(
        select(Chunk).where(Chunk.document_id == document.id)
    ).scalars().all()
    assert [chunk.id for chunk in after_duplicate] == chunk_ids


def test_revoked_worker_lease_cannot_refresh_or_change_retried_job(db: Session) -> None:
    course = add_course(db, "Course 1")
    document = Document(
        course_id=course.id,
        uploaded_by=uuid4(),
        filename=SAMPLE_MD.name,
        storage_path=str(SAMPLE_MD),
    )
    old_token = uuid4()
    db.add(document)
    db.flush()
    job = IngestJob(
        document_id=document.id,
        status=IngestJobStatus.running.value,
        run_token=old_token,
        updated_at=datetime.now(UTC) - timedelta(hours=1),
    )
    db.add(job)
    db.commit()

    assert retry_ingest_job(db, job.id, datetime.now(UTC) - timedelta(minutes=15))
    with pytest.raises(IngestLeaseLost, match="lease lost"):
        _touch_running_job(db, job.id, old_token)

    db.expire_all()
    stored_job = db.get(IngestJob, job.id)
    assert stored_job is not None
    assert stored_job.status == IngestJobStatus.pending.value
    assert stored_job.run_token is None


def test_active_ingest_job_cannot_be_retried(client: TestClient, db: Session) -> None:
    teacher = add_user(db, "t1", "pass", UserRole.teacher)
    course = add_course(db, "Course 1")
    add_member(db, course, teacher, CourseMemberRole.teacher)
    document = Document(
        course_id=course.id,
        uploaded_by=teacher.id,
        filename=SAMPLE_MD.name,
        storage_path=str(SAMPLE_MD),
    )
    db.add(document)
    db.flush()
    db.add(
        IngestJob(
            document_id=document.id,
            status=IngestJobStatus.running.value,
            run_token=uuid4(),
            updated_at=datetime.now(UTC),
        )
    )
    db.commit()
    login_as(client, "t1", "pass")

    response = client.post(f"/courses/{course.id}/documents/{document.id}/retry")

    assert response.status_code == 409


def test_student_cannot_retry_ingest_job(client: TestClient, db: Session) -> None:
    teacher = add_user(db, "t1", "pass", UserRole.teacher)
    student = add_user(db, "alice", "pass", UserRole.student)
    course = add_course(db, "Course 1")
    add_member(db, course, teacher, CourseMemberRole.teacher)
    add_member(db, course, student, CourseMemberRole.student)
    document = Document(
        course_id=course.id,
        uploaded_by=teacher.id,
        filename=SAMPLE_MD.name,
        storage_path=str(SAMPLE_MD),
    )
    db.add(document)
    db.flush()
    db.add(IngestJob(document_id=document.id, status=IngestJobStatus.failed.value))
    db.commit()
    login_as(client, "alice", "pass")

    response = client.post(f"/courses/{course.id}/documents/{document.id}/retry")

    assert response.status_code == 403


@pytest.mark.parametrize("filename", ["synthetic.md", "synthetic.txt"])
def test_ingest_source_spans_match_sanitized_extracted_text_and_embedding_input(
    client, db, filename,
) -> None:
    body = " \r\n# 课程标题\r\n\r\n" + "中文🙂e\u0301重复正文\x00。" * 220

    class CapturingEmbedder(FakeEmbedder):
        def __init__(self):
            super().__init__()
            self.texts = []

        def embed(self, texts):
            self.texts.extend(texts)
            return super().embed(texts)

    embedder = CapturingEmbedder()
    app.dependency_overrides[get_embedder] = lambda: embedder
    doc_id = upload_as_teacher(client, db, filename, body.encode())
    document = db.get(Document, doc_id)
    units = parse_document(filename, body.encode())
    sanitized = units[0][1].replace("\x00", "")
    chunks = db.execute(select(Chunk).where(Chunk.document_id == doc_id)).scalars().all()
    chunks.sort(key=lambda chunk: chunk.source_metadata["chunk_index"])
    assert document.chunking_version == CHUNKING_VERSION
    assert document.chunk_count == len(chunks) > 1
    assert [chunk.text for chunk in chunks] == embedder.texts
    for chunk in chunks:
        meta = chunk.source_metadata
        assert sanitized[meta["char_start"]:meta["char_end"]] == chunk.text
        assert meta["unit_index"] == 1
        assert chunk.page is None
    assert document.extracted_char_count == len(sanitized)
    detail = client.get(f"/courses/{document.course_id}/documents/{doc_id}").json()
    assert detail["chunking_version"] == CHUNKING_VERSION
    assert {"text", "source_metadata", "storage_path"}.isdisjoint(detail)


@pytest.mark.parametrize("fail_on_call", [None, 2])
def test_reingest_replaces_source_metadata_and_version_only_after_complete_success(
    db, tmp_path, fail_on_call,
) -> None:
    course = add_course(db, "Reingest provenance")
    source = tmp_path / "synthetic.txt"
    source.write_text("new content " * 500, encoding="utf-8")
    document = Document(
        course_id=course.id, uploaded_by=uuid4(), filename=source.name, storage_path=str(source),
        chunking_version="legacy-v0", chunk_count=1,
    )
    db.add(document)
    db.flush()
    old = Chunk(
        course_id=course.id, document_id=document.id, page=1, text="old complete index",
        embedding=FakeEmbedder().embed(["old complete index"])[0],
        embedding_model=FakeEmbedder().index_identity, source_metadata=None,
    )
    job = IngestJob(document_id=document.id, status=IngestJobStatus.pending.value)
    db.add_all([old, job])
    db.commit()
    old_id = old.id
    process_ingest_job(db, job.id, RecordingEmbedder(fail_on_call=fail_on_call), batch_size=1)
    db.expire_all()
    chunks = db.execute(select(Chunk).where(Chunk.document_id == document.id)).scalars().all()
    stored = db.get(Document, document.id)
    if fail_on_call:
        assert db.get(IngestJob, job.id).status == IngestJobStatus.failed.value
        assert stored.chunking_version == "legacy-v0"
        assert stored.chunk_count == 1
        assert [chunk.id for chunk in chunks] == [old_id]
        assert chunks[0].source_metadata is None
    else:
        assert db.get(IngestJob, job.id).status == IngestJobStatus.succeeded.value
        assert stored.chunking_version == CHUNKING_VERSION
        assert old_id not in {chunk.id for chunk in chunks}
        assert stored.chunk_count == len(chunks) > 1
        assert all(chunk.source_metadata is not None for chunk in chunks)
        assert all(chunk.page is None for chunk in chunks)


@pytest.mark.parametrize(
    ("mode", "expected_error"),
    [("partial", None), ("sparse", None), ("unreadable", "unreadable PDF"),
     ("empty", "no extractable text"), ("failed_report", "incomplete PDF parse"),
     ("success", None)],
)
def test_pdf_quality_gates_atomic_reingest(
    db, tmp_path, monkeypatch, mode, expected_error,
) -> None:
    source = tmp_path / "fake.pdf"
    text = "short" if mode == "sparse" else ("new content " * 100).strip()
    data = make_pdf(text, "") if mode == "partial" else make_pdf(text)
    if mode == "empty":
        data = make_pdf("")
    elif mode == "unreadable":
        data = b"not a PDF"
    source.write_bytes(data)
    course = add_course(db, "PDF quality")
    document = Document(
        course_id=course.id, uploaded_by=uuid4(), filename=source.name,
        storage_path=str(source), chunking_version="legacy-v0", chunk_count=1,
        source_unit_count=1, extracted_char_count=18, quality_warning=LOW_EXTRACTED_TEXT,
    )
    db.add(document)
    db.flush()
    old = Chunk(
        course_id=course.id, document_id=document.id, page=1, text="old complete index",
        embedding=FakeEmbedder().embed(["old complete index"])[0],
        embedding_model=FakeEmbedder().index_identity,
        source_metadata={"kind": "pdf_page", "unit_index": 1},
    )
    job = IngestJob(document_id=document.id, status=IngestJobStatus.pending.value)
    db.add_all([old, job])
    db.commit()
    old_id = old.id
    if mode == "failed_report":
        from dataclasses import replace

        from app.services.parse import parse_document_with_quality
        from app.services.parse_quality import ParseStatus

        pages, report = parse_document_with_quality(source.name, data)
        monkeypatch.setattr(
            "app.services.ingest.parse_document_with_quality",
            lambda *args, **kwargs: (pages, replace(report, status=ParseStatus.FAILED)),
        )
    embedder = RecordingEmbedder()
    process_ingest_job(db, job.id, embedder)
    db.expire_all()
    stored = db.get(Document, document.id)
    stored_job = db.get(IngestJob, job.id)
    chunks = db.execute(select(Chunk).where(Chunk.document_id == document.id)).scalars().all()
    if expected_error:
        assert stored_job.status == IngestJobStatus.failed.value
        assert stored_job.error_message == expected_error
        assert stored_job.run_token is None
        assert not embedder.batch_sizes
        assert [chunk.id for chunk in chunks] == [old_id]
        assert chunks[0].text == "old complete index"
        assert chunks[0].source_metadata == {"kind": "pdf_page", "unit_index": 1}
        assert stored.chunking_version == "legacy-v0"
        assert (stored.source_unit_count, stored.extracted_char_count, stored.chunk_count) == (1, 18, 1)
        assert stored.quality_warning == LOW_EXTRACTED_TEXT
    else:
        assert stored_job.status == IngestJobStatus.succeeded.value
        assert old_id not in {chunk.id for chunk in chunks}
        assert stored.chunk_count == len(chunks) >= 1
        assert stored.source_unit_count == 1
        assert stored.extracted_char_count == len(text)
        assert stored.chunking_version == CHUNKING_VERSION
        assert stored_job.error_message is None
        assert stored.quality_warning == (
            COVERAGE_INCOMPLETE if mode in {"partial", "sparse"} else None
        )
        for chunk in chunks:
            assert chunk.page == chunk.source_metadata["unit_index"] == 1
            assert "parser_version" in chunk.source_metadata
            assert len(chunk.source_metadata["structured_hash"]) == 64
            assert all(isinstance(value, (str, int)) for value in chunk.source_metadata.values())
