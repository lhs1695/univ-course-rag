"""Synthetic parser/order regressions for the R5 splitter inputs."""

import builtins
import hashlib
import io

import fitz
import pytest
from app.services.parse import (
    ParseError,
    ParseLimits,
    parse_document,
    parse_document_with_quality,
)
from docx import Document
from pptx import Presentation


def test_word_interleaves_tables_and_paragraphs_in_document_order() -> None:
    document = Document()
    document.add_paragraph("First heading", style="Heading 1")
    first = document.add_table(rows=1, cols=3)
    first.cell(0, 0).text = "First table left"
    first.cell(0, 2).text = "First table right"
    document.add_paragraph("Between the tables")
    second = document.add_table(rows=1, cols=1)
    second.cell(0, 0).text = "Second table"
    document.add_paragraph("Final paragraph")
    output = io.BytesIO()
    document.save(output)

    units = parse_document("synthetic.docx", output.getvalue())

    body = units[0][1]
    assert body == (
        "First heading\n\nFirst table left\t\tFirst table right\n\n"
        "Between the tables\n\nSecond table\n\nFinal paragraph"
    )


def make_pdf(*texts: str) -> bytes:
    with fitz.open() as document:
        for text in texts:
            page = document.new_page()
            if text:
                page.insert_textbox(fitz.Rect(30, 30, 570, 800), text)
        return document.tobytes()


def test_pdf_skips_empty_page_without_renumbering_later_sources() -> None:
    data = make_pdf("First page", "", "Third page")
    units, report = parse_document_with_quality("synthetic.pdf", data)
    assert units == [(1, "First page"), (3, "Third page")]
    assert parse_document("synthetic.pdf", data) == units
    assert report.total_pages == 3
    assert [p.number for p in report.pages] == [1, 2, 3]
    assert report.page_numbers_continuous
    assert report.empty_pages == (2,)
    assert report.status.value == "partial"
    assert report.input_sha256 == hashlib.sha256(data).hexdigest()
    assert report.structured_hash == parse_document_with_quality("x.pdf", data)[1].structured_hash
    assert report.structured_hash != parse_document_with_quality(
        "x.pdf", make_pdf("Changed", "", "Third page"),
    )[1].structured_hash


@pytest.mark.parametrize("size, status", [(49, "partial"), (50, "succeeded")])
def test_pdf_sparse_threshold(size, status) -> None:
    _, report = parse_document_with_quality("x.pdf", make_pdf("x" * size))
    assert report.status.value == status
    assert report.pages[0].char_length == size
    assert report.pages[0].sparse is (size < 50)


@pytest.mark.parametrize("control", ["\x00", "\x01", "\x85"])
def test_pdf_abnormal_controls_are_coverage_warnings(monkeypatch, control) -> None:
    real_open = fitz.open
    class ExtractedPage:
        number = 0
        def get_text(self, mode):
            assert mode == "text"
            return "a" * 50 + control + "\t\n\r"
    class ExtractedDocument:
        page_count = 1
        def __iter__(self):
            return iter([ExtractedPage()])
        def close(self):
            pass
    data = make_pdf("a" * 50)
    monkeypatch.setattr(fitz, "open", lambda **kwargs: ExtractedDocument())
    _, report = parse_document_with_quality("x.pdf", data)
    assert report.status.value == "partial"
    assert report.pages[0].abnormal_chars
    assert "abnormal_characters" in report.pages[0].notes
    monkeypatch.setattr(fitz, "open", real_open)


@pytest.mark.parametrize("failure", ["count", "number", "execution"])
def test_pdf_hard_mapping_and_execution_failures(monkeypatch, failure) -> None:
    class BrokenPage:
        number = 3 if failure == "number" else 0
        def get_text(self, mode):
            if failure == "execution":
                raise RuntimeError("private error")
            return "readable text"
    class BrokenDocument:
        page_count = 2 if failure == "count" else 1
        closed = False
        def __iter__(self):
            return iter([BrokenPage()])
        def close(self):
            self.closed = True
    doc = BrokenDocument()
    monkeypatch.setattr(fitz, "open", lambda **kwargs: doc)
    message = "unreadable PDF" if failure == "execution" else "incomplete PDF parse"
    with pytest.raises(ParseError, match=message):
        parse_document("x.pdf", b"pdf")
    assert doc.closed


@pytest.mark.parametrize("limits", [ParseLimits(max_pages=1), ParseLimits(max_text_chars=5)])
def test_native_pdf_resource_limits(limits) -> None:
    with pytest.raises(ParseError, match="extracted text exceeds resource limit"):
        parse_document("x.pdf", make_pdf("first page", "second page"), limits=limits)


def test_default_pdf_does_not_import_docling(monkeypatch) -> None:
    data = make_pdf("Native text")
    original = builtins.__import__
    def guarded(name, *args, **kwargs):
        if name.startswith("docling"):
            raise AssertionError("default path imported Docling")
        return original(name, *args, **kwargs)
    monkeypatch.setattr(builtins, "__import__", guarded)
    assert parse_document("x.pdf", data) == [(1, "Native text")]


def test_missing_pymupdf_uses_safe_message(monkeypatch) -> None:
    original = builtins.__import__
    def guarded(name, *args, **kwargs):
        if name == "fitz":
            raise ImportError("missing")
        return original(name, *args, **kwargs)
    monkeypatch.setattr(builtins, "__import__", guarded)
    with pytest.raises(ParseError, match="PDF parser is not installed"):
        parse_document("x.pdf", b"pdf")


def test_non_pdf_quality_is_optional() -> None:
    assert parse_document_with_quality("note.md", b"hello") == ([(1, "hello")], None)


def test_slides_skip_empty_slide_without_renumbering_later_sources() -> None:
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    presentation.slides.add_slide(blank).shapes.add_textbox(0, 0, 100, 100).text = "First slide"
    presentation.slides.add_slide(blank)
    presentation.slides.add_slide(blank).shapes.add_textbox(0, 0, 100, 100).text = "Third slide"
    output = io.BytesIO()
    presentation.save(output)
    units = parse_document("synthetic.pptx", output.getvalue())
    assert [page for page, _text in units] == [1, 3]
    assert [text for _page, text in units] == ["First slide", "Third slide"]
