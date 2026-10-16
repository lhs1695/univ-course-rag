from __future__ import annotations

import io
from dataclasses import dataclass
from zipfile import BadZipFile, ZipFile

from app.config import Settings
from app.services.docling_adapter import (
    DoclingAdapterError,
    ParseQualityReport,
    PdfConverter,
    build_parse_quality_report,
    flatten_pages,
    parse_pdf_document,
)


class ParseError(Exception):
    pass


_UNSUPPORTED = "unsupported file type; accepted: PDF, PPTX, DOCX, Markdown, or plain text"
_OOXML_RESOURCE_LIMIT = "OOXML resource limit exceeded"
_TEXT_RESOURCE_LIMIT = "extracted text exceeds resource limit"


@dataclass(frozen=True)
class ParseLimits:
    max_zip_entries: int = 2_000
    max_uncompressed_bytes: int = 50 * 1024 * 1024
    max_entry_bytes: int = 10 * 1024 * 1024
    max_compression_ratio: float = 100.0
    max_pages: int = 500
    max_text_chars: int = 2_000_000


DEFAULT_PARSE_LIMITS = ParseLimits()


def parse_limits_from_settings(settings: Settings) -> ParseLimits:
    return ParseLimits(
        max_zip_entries=settings.ooxml_max_zip_entries,
        max_uncompressed_bytes=settings.ooxml_max_uncompressed_bytes,
        max_entry_bytes=settings.ooxml_max_entry_bytes,
        max_compression_ratio=settings.ooxml_max_compression_ratio,
        max_pages=settings.parse_max_pages,
        max_text_chars=settings.parse_max_text_chars,
    )


def parse_document(
    filename: str,
    data: bytes,
    *,
    limits: ParseLimits = DEFAULT_PARSE_LIMITS,
    pdf_converter: PdfConverter | None = None,
) -> list[tuple[int | None, str]]:
    """Return (page, text) pairs. Page is 1-based for PDF/PPTX, 1 otherwise."""
    pages, _quality = parse_document_with_quality(
        filename, data, limits=limits, pdf_converter=pdf_converter,
    )
    return pages


def parse_document_with_quality(
    filename: str,
    data: bytes,
    *,
    limits: ParseLimits = DEFAULT_PARSE_LIMITS,
    pdf_converter: PdfConverter | None = None,
) -> tuple[list[tuple[int | None, str]], ParseQualityReport | None]:
    """PDF quality is returned for inspection; ingest must gate on its status."""
    name = filename.lower()
    if name.endswith(".pdf"):
        try:
            document = parse_pdf_document(
                data, converter=pdf_converter, max_pages=limits.max_pages,
                max_text_chars=limits.max_text_chars,
            )
        except DoclingAdapterError as exc:
            raise ParseError(str(exc)) from exc
        return flatten_pages(document), build_parse_quality_report(document)
    if name.endswith(".pptx"):
        _validate_ooxml_archive(data, limits)
        return _parse_ppt(data, limits), None
    if name.endswith(".docx"):
        _validate_ooxml_archive(data, limits)
        return _parse_word(data, limits), None
    if name.endswith((".md", ".markdown", ".txt")):
        return _parse_text(data, limits), None
    raise ParseError(_UNSUPPORTED)


def _parse_text(data: bytes, limits: ParseLimits) -> list[tuple[int | None, str]]:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ParseError("text file is not valid UTF-8") from exc
    text = text.replace("\r\n", "\n").strip()
    _validate_text_size(len(text), limits)
    if not text:
        raise ParseError("no extractable text")
    return [(1, text)]


def _parse_ppt(data: bytes, limits: ParseLimits) -> list[tuple[int | None, str]]:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        raise ParseError("PPT parser is not installed") from exc
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ParseError("unreadable PPT") from exc

    pages: list[tuple[int | None, str]] = []
    extracted_chars = 0
    for index, slide in enumerate(presentation.slides, start=1):
        if index > limits.max_pages:
            raise ParseError(_TEXT_RESOURCE_LIMIT)
        parts: list[str] = []
        for shape in slide.shapes:
            parts.extend(_ppt_shape_text(shape, MSO_SHAPE_TYPE))
        page_text = "\n".join(parts).replace("\r\n", "\n").strip()
        if page_text:
            extracted_chars += len(page_text)
            _validate_text_size(extracted_chars, limits)
            pages.append((index, page_text))
    if not pages:
        raise ParseError("no extractable text")
    return pages


def _ppt_shape_text(shape: object, mso_shape_type: object) -> list[str]:
    parts: list[str] = []
    shape_type = getattr(shape, "shape_type", None)
    if shape_type == mso_shape_type.GROUP:
        for inner in shape.shapes:
            parts.extend(_ppt_shape_text(inner, mso_shape_type))
        return parts
    if getattr(shape, "has_text_frame", False):
        text = shape.text_frame.text.replace("\r\n", "\n").strip()
        if text:
            parts.append(text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [cell.text.replace("\r\n", "\n").strip() for cell in row.cells]
            line = "\t".join(cell for cell in cells if cell)
            if line:
                parts.append(line)
    return parts


def _parse_word(data: bytes, limits: ParseLimits) -> list[tuple[int | None, str]]:
    try:
        from docx import Document
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise ParseError("Word parser is not installed") from exc
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ParseError("unreadable Word") from exc

    parts: list[str] = []
    # Public python-docx 1.2.0 API preserves paragraph/table document order.
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            text = block.text.replace("\r\n", "\n").strip()
            if text:
                parts.append(text)
        else:
            rows = []
            for row in block.rows:
                cells = [cell.text.replace("\r\n", "\n").strip() for cell in row.cells]
                line = "\t".join(cells)
                if line.strip():
                    rows.append(line)
            if rows:
                parts.append("\n".join(rows))
    page_text = "\n\n".join(parts).strip()
    _validate_text_size(len(page_text), limits)
    if not page_text:
        raise ParseError("no extractable text")
    return [(1, page_text)]


def _validate_ooxml_archive(data: bytes, limits: ParseLimits) -> None:
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
    except (BadZipFile, OSError) as exc:
        raise ParseError("unreadable OOXML") from exc
    if len(entries) > limits.max_zip_entries:
        raise ParseError(_OOXML_RESOURCE_LIMIT)

    total_uncompressed = 0
    for entry in entries:
        if entry.file_size > limits.max_entry_bytes:
            raise ParseError(_OOXML_RESOURCE_LIMIT)
        total_uncompressed += entry.file_size
        if total_uncompressed > limits.max_uncompressed_bytes:
            raise ParseError(_OOXML_RESOURCE_LIMIT)
        compressed_size = max(entry.compress_size, 1)
        if entry.file_size / compressed_size > limits.max_compression_ratio:
            raise ParseError(_OOXML_RESOURCE_LIMIT)


def _validate_text_size(text_size: int, limits: ParseLimits) -> None:
    if text_size > limits.max_text_chars:
        raise ParseError(_TEXT_RESOURCE_LIMIT)
