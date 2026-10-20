from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path.cwd()
BASE = ROOT / "build/midterm-presentation"
OUT = BASE / "output"
FINALIZER = BASE / "finalizer"
OUT.mkdir(parents=True, exist_ok=True)
FINALIZER.mkdir(parents=True, exist_ok=True)
PPTX_PATH = OUT / "final-v5.pptx"
VALIDATION_PATH = FINALIZER / "final-v5.validation.json"
NOTES_PATH = OUT / "speaker-notes-v5.md"
HANDOFF_PPT = ROOT / "中期答辩材料/汇报PPT/中期汇报-v5.pptx"
HANDOFF_NOTES = ROOT / "中期答辩材料/汇报PPT/讲稿-v5.md"
HANDOFF_SHOTS = ROOT / "中期答辩材料/彩排截图"

W, H = 13.333, 7.5
FONT = "Microsoft YaHei"
NAVY = "163247"
TEAL = "247D78"
TEAL_LIGHT = "E4F2F0"
AMBER = "D39A37"
AMBER_LIGHT = "FBF0D8"
RED = "B84242"
RED_LIGHT = "F9E7E5"
GREEN = "26734E"
GREEN_LIGHT = "E4F1E9"
INK = "263642"
MUTED = "61727D"
LINE = "D6E0E5"
BG = "F5F8FA"
WHITE = "FFFFFF"

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return shape


def add_box(slide, x, y, w, h, fill=WHITE, stroke=LINE, radius=True, line_width=1.1):
    geometry = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(geometry, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(stroke)
    shape.line.width = Pt(line_width)
    return shape


def add_card(slide, x, y, w, h, title, body, fill=WHITE, stroke=LINE, title_color=NAVY, body_color=INK, title_size=19, body_size=15):
    shape = add_box(slide, x, y, w, h, fill, stroke)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.16)
    tf.margin_top = Inches(0.10)
    tf.margin_bottom = Inches(0.08)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(6)
    r = p.add_run()
    r.text = title
    r.font.name = FONT
    r.font.size = Pt(title_size)
    r.font.bold = True
    r.font.color.rgb = rgb(title_color)
    for line in body.split("\n"):
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(2)
        r = p.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(body_size)
        r.font.color.rgb = rgb(body_color)
    return shape


def add_label(slide, text, x, y, w, fill=TEAL_LIGHT, color=TEAL):
    add_box(slide, x, y, w, 0.38, fill, fill, radius=True, line_width=0.5)
    add_text(slide, text, x + 0.05, y + 0.015, w - 0.10, 0.34, 14, color, True, PP_ALIGN.CENTER)


def add_line(slide, x1, y1, x2, y2, color=TEAL, width=2.0):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    return shape


def add_header(slide, title, subtitle, page, total=10):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)
    add_text(slide, title, 0.62, 0.34, 12.0, 0.52, 27, NAVY, True)
    add_text(slide, subtitle, 0.66, 0.91, 11.9, 0.40, 14, MUTED)
    add_box(slide, 0.66, 1.40, 0.72, 0.045, AMBER, AMBER, radius=False, line_width=0.2)
    add_text(slide, "高校课程知识库智能问答 · 中期进展（参考稿）", 0.62, 7.12, 9.5, 0.20, 9, MUTED)
    add_text(slide, f"{page:02d} / {total:02d}", 11.7, 7.10, 0.95, 0.22, 9, MUTED, align=PP_ALIGN.RIGHT)


def add_slide(title, subtitle, page, total=10):
    slide = prs.slides.add_slide(blank)
    add_header(slide, title, subtitle, page, total)
    return slide


def add_stat(slide, x, y, w, value, label, color=TEAL):
    add_box(slide, x, y, w, 1.12, WHITE, LINE)
    add_text(slide, value, x + 0.08, y + 0.10, w - 0.16, 0.50, 26, color, True, PP_ALIGN.CENTER)
    add_text(slide, label, x + 0.08, y + 0.64, w - 0.16, 0.30, 13, MUTED, align=PP_ALIGN.CENTER)


def add_image_contain(slide, path: Path, x, y, w, h):
    with Image.open(path) as im:
        ratio = im.width / im.height
    out_w = min(w, h * ratio)
    out_h = out_w / ratio
    left = x + (w - out_w) / 2
    top = y + (h - out_h) / 2
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(out_w), height=Inches(out_h))


TOTAL = 10

# 1. Cover
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = rgb(NAVY)
add_box(slide, 0, 0, 6.6, H, NAVY, NAVY, radius=False, line_width=0.1)
add_box(slide, 7.05, 1.35, 5.55, 4.95, "1D4053", "416173", radius=True, line_width=1.0)
add_card(slide, 7.40, 1.80, 2.20, 1.15, "课程资料", "4 门课 · 27 份 PDF", TEAL_LIGHT, TEAL, TEAL, INK, 17, 13)
add_card(slide, 10.00, 1.80, 2.20, 1.15, "演示检索", "默认 Dense", WHITE, WHITE, NAVY, INK, 17, 13)
add_line(slide, 9.60, 2.38, 9.95, 2.38, AMBER, 2.0)
add_card(slide, 7.40, 3.30, 2.20, 1.15, "生成", "DeepSeek\ndeepseek-flash", AMBER_LIGHT, AMBER, AMBER, INK, 16, 12)
add_card(slide, 10.00, 3.30, 2.20, 1.15, "核对", "引用回原 PDF", WHITE, WHITE, NAVY, INK, 17, 13)
add_card(slide, 8.35, 4.75, 2.95, 1.10, "教学反馈", "高频问句 · 依据不足线索", GREEN_LIGHT, GREEN, GREEN, INK, 16, 12)
add_text(slide, "检索命中 ≠ 答案正确", 7.40, 6.05, 4.85, 0.32, 16, "D8E7EE", True, PP_ALIGN.CENTER)
add_text(slide, "高校课程知识库\n智能问答", 0.78, 1.55, 5.6, 1.55, 36, WHITE, True)
add_box(slide, 0.82, 3.40, 0.9, 0.06, AMBER, AMBER, radius=False, line_width=0.1)
add_text(slide, "中期进展汇报", 0.82, 3.70, 5.2, 0.50, 24, "EAF2F6", True)
add_text(slide, "有依据、可核对的课程问答原型", 0.84, 4.35, 5.3, 0.36, 16, "D3E1E8")
add_text(slide, "参考稿，队友可改；终稿可由答辩同学重做", 0.84, 5.85, 5.4, 0.36, 14, AMBER, True)
add_text(slide, "中文四课文本型 RAG · 2026.10", 0.84, 6.55, 5.5, 0.34, 14, "D3E1E8")

# 2. Problem
slide = add_slide("问题背景与中期目标", "课件分散难定位；通用模型流畅≠课程依据；需要可核对的问答闭环", 2, total=TOTAL)
add_text(slide, "实际问题", 0.78, 1.70, 2.0, 0.36, 18, TEAL, True)
add_card(slide, 0.75, 2.12, 5.85, 3.55,
         "为什么要做",
         "• 课程知识分散在多份讲义/课件中\n• 自然语言问题难以靠文件名检索定位\n• 通用大模型回答流畅，但未必对应当前课程口径\n• 学生需要原文核对入口；教师需要问题分布线索\n• 多课程场景必须守住访问边界",
         WHITE, LINE, NAVY, INK, 18, 15)
add_text(slide, "中期目标", 7.00, 1.70, 2.0, 0.36, 18, AMBER, True)
add_card(slide, 6.95, 2.12, 5.55, 3.55,
         "本阶段交付什么",
         "• 本机可演示的文本型 RAG 原型\n• 中文四课受控资料：4 / 27 / 2,501\n• 学生能提问、看引用、回原 PDF 页\n• 依据不足与系统失败分开呈现\n• 教师侧基础教学反馈（统计/待核对）\n• 不宣称生产部署或通用多模态",
         AMBER_LIGHT, AMBER, AMBER, INK, 18, 15)
add_text(slide, "核心价值：有依据、可核对的课程问答；教学反馈是辅助闭环。", 0.85, 6.05, 11.6, 0.36, 17, NAVY, True, PP_ALIGN.CENTER)

# 3. Closed loop
slide = add_slide("业务闭环", "课程资料管理 → 授权检索 → 证据问答 → 来源核对 → 教学反馈", 3, total=TOTAL)
steps = [
    (0.55, "① 资料管理", "上传解析入库\n页级质量门禁", TEAL_LIGHT, TEAL),
    (3.05, "② 授权检索", "课程成员边界\n课内候选召回", WHITE, TEAL),
    (5.55, "③ 证据问答", "受控上下文生成\n有据 / 拒答 / 失败", AMBER_LIGHT, AMBER),
    (8.05, "④ 来源核对", "引用摘录\n打开原 PDF 页", WHITE, AMBER),
    (10.55, "⑤ 教学反馈", "高频问句\n依据不足线索", GREEN_LIGHT, GREEN),
]
for i, (x, title, body, fill, stroke) in enumerate(steps):
    add_card(slide, x, 1.85, 2.25, 2.55, title, body, fill, stroke, stroke if fill != WHITE else NAVY, INK, 17, 14)
    if i < len(steps) - 1:
        add_line(slide, x + 2.28, 3.05, steps[i + 1][0] - 0.08, 3.05, AMBER, 2.0)
add_box(slide, 0.70, 4.75, 11.9, 1.35, WHITE, LINE)
add_text(slide, "三角色分工", 0.95, 4.95, 2.0, 0.34, 17, TEAL, True)
add_text(slide, "管理员：账号与课程、成员授权　｜　教师：本课资料状态与基础反馈　｜　学生：我的课程 → 提问 → 查看引用 → 回原页", 0.95, 5.40, 11.3, 0.45, 16, INK)

# 4. Dual narrative highlights
slide = add_slide("技术亮点 · 双叙事", "教学场景闭环 + 检索/生成技术；中期只讲能演示、有证据的能力", 4, total=TOTAL)
add_card(slide, 0.55, 1.75, 6.05, 4.35,
         "A. 教学场景闭环",
         "• 课界权限贯穿检索—生成—打开原文件\n• 证据不足可拒答，不编造课外步骤\n• 引用带回原 PDF 页（#page=）\n• 三角色分工清晰，不是单一聊天窗\n• 教师问答反馈：高频问句与依据不足线索",
         TEAL_LIGHT, TEAL, TEAL, INK, 19, 15)
add_card(slide, 6.85, 1.75, 5.85, 4.35,
         "B. 检索 / 生成技术",
         "• 页级解析 + 质量门禁（PyMuPDF）\n• 向量检索 BGE-M3 + 证据约束生成\n• DeepSeek（deepseek-flash）\n• dense / hybrid / rerank 已实现可开关\n• 中期主演示：dense，rerank 关\n• Hybrid 对照实验放后期，不作主结论",
         AMBER_LIGHT, AMBER, AMBER, INK, 19, 15)

# 5. Architecture
slide = add_slide("系统链路与技术栈", "复用成熟组件；本地实现课程边界、来源映射与失败状态", 5, total=TOTAL)
nodes = [
    (0.55, "课程 PDF", "原始文件\n课程隔离", TEAL_LIGHT, TEAL),
    (2.70, "解析质检", "PyMuPDF\n页级门禁", WHITE, TEAL),
    (4.85, "切分索引", "语义切分\nBGE-M3", WHITE, TEAL),
    (7.00, "检索", "默认 Dense\n可选 Hybrid", WHITE, TEAL),
    (9.15, "生成", "DeepSeek\ndeepseek-flash", AMBER_LIGHT, AMBER),
    (11.05, "学生答案", "回答/引用\n依据不足", NAVY, NAVY),
]
for i, (x, title, body, fill, stroke) in enumerate(nodes):
    add_box(slide, x, 1.85, 1.85, 1.55, fill, stroke)
    add_text(slide, title, x + 0.08, 2.00, 1.70, 0.38, 15, WHITE if fill == NAVY else NAVY, True, PP_ALIGN.CENTER)
    add_text(slide, body, x + 0.08, 2.45, 1.70, 0.70, 12, WHITE if fill == NAVY else MUTED, align=PP_ALIGN.CENTER)
    if i < len(nodes) - 1:
        add_line(slide, x + 1.88, 2.60, nodes[i + 1][0] - 0.05, 2.60, AMBER if i >= 3 else TEAL, 1.8)
add_box(slide, 0.70, 3.75, 11.9, 2.35, WHITE, LINE)
add_text(slide, "应用层栈", 0.95, 3.95, 2.0, 0.34, 17, TEAL, True)
add_text(slide, "FastAPI + PostgreSQL/pgvector + Vue 3 · Cookie Session · admin / teacher / student", 0.95, 4.40, 11.3, 0.36, 15, INK)
add_text(slide, "Embedding：SiliconFlow BAAI/bge-m3　｜　生成：DeepSeek（deepseek-flash，关闭 thinking）", 0.95, 4.85, 11.3, 0.36, 15, INK)
add_text(slide, "中期演示默认 Dense；Hybrid/BM25/rerank 已实现可开关，不作中期主结论。", 0.95, 5.35, 11.3, 0.40, 15, AMBER, True)

# 6. Midterm evidence
slide = add_slide("中文四课中期证据", "正式演示库 univ_course_rag_cn_v2；弱规则验收 + 彩排 PASS，不是人工事实正确率", 6, total=TOTAL)
add_stat(slide, 0.70, 1.75, 2.70, "4", "门课程")
add_stat(slide, 3.60, 1.75, 2.70, "27", "份 PDF", AMBER)
add_stat(slide, 6.50, 1.75, 2.70, "2501", "chunks", GREEN)
add_stat(slide, 9.40, 1.75, 2.90, "cn_v2", "演示库", NAVY)
add_card(slide, 0.70, 3.20, 5.90, 2.55,
         "课程与验收口径",
         "线性代数 · 概率论与统计学\n人力资源管理 · 教育学原理\n弱规则：可答题预期文档/页 12/12\n预期拒答 4/4 · 阈值 0.45（未校准）\n题集 AI 预标，未经独立人工复核",
         WHITE, LINE, TEAL, INK, 18, 14)
add_card(slide, 6.85, 3.20, 5.75, 2.55,
         "演示检索与模型",
         "默认 Dense + 可选查询扩展\nrerank 演示关闭\n生成：DeepSeek（deepseek-flash）\nEmbedding：BGE-M3\n证据目录：中期答辩材料/彩排截图",
         AMBER_LIGHT, AMBER, AMBER, INK, 18, 14)

# 7. Demo path
slide = add_slide("演示路径：固定问句与引用回页", "线性代数 LA-01 有引用；LA-R 依据不足；截图见中期答辩材料/彩排截图", 7, total=TOTAL)
screens = BASE / "assets-v5"
screens.mkdir(parents=True, exist_ok=True)
left_source = HANDOFF_SHOTS / "03-学生-线代作答与引用.png"
right_source = HANDOFF_SHOTS / "01-学生-课程书架.png"
if not left_source.exists():
    left_source = HANDOFF_SHOTS / "05-学生-拒答.png"
left_img = screens / "student-ask-focus.png"
right_img = screens / "student-shelf-focus.png"
for source, destination in ((left_source, left_img), (right_source, right_img)):
    if not source.exists():
        raise FileNotFoundError(f"Missing screenshot: {source}")
    with Image.open(source) as original:
        top = 0
        bottom = original.height
        if original.height > 1100:
            top = min(40, original.height - 1000)
            bottom = min(original.height, top + 1000)
        original.crop((0, top, original.width, bottom)).convert("RGB").save(destination, quality=90)
add_box(slide, 0.65, 1.70, 5.55, 3.35, WHITE, LINE)
add_image_contain(slide, left_img, 0.78, 1.80, 5.30, 2.85)
add_text(slide, "LA-01：回答与引用", 0.80, 4.70, 5.20, 0.25, 13, NAVY, True, PP_ALIGN.CENTER)
add_box(slide, 6.95, 1.70, 5.55, 3.35, WHITE, LINE)
add_image_contain(slide, right_img, 7.08, 1.80, 5.30, 2.85)
add_text(slide, "学生书架：中文四课入口", 7.10, 4.70, 5.20, 0.25, 13, NAVY, True, PP_ALIGN.CENTER)
add_label(slide, "LA-01 有引用", 0.70, 5.20, 2.90, GREEN_LIGHT, GREEN)
add_label(slide, "LA-R 依据不足", 3.80, 5.20, 2.90, AMBER_LIGHT, AMBER)
add_label(slide, "打开 #page=", 6.90, 5.20, 2.70, TEAL_LIGHT, TEAL)
add_label(slide, "三角色可走通", 9.80, 5.20, 2.70, GREEN_LIGHT, GREEN)
add_text(slide, "固定问句与顺序见 docs/中期演示运行卡.md；证据截图：中期答辩材料/彩排截图/", 0.80, 5.90, 11.5, 0.32, 14, MUTED, align=PP_ALIGN.CENTER)

# 8. Teaching feedback
slide = add_slide("教学反馈：辅助闭环", "把重复提问与依据不足整理成教师可核对的线索，不直接推断学生能力", 8, total=TOTAL)
add_card(slide, 0.70, 1.80, 5.90, 4.15,
         "中期已交付",
         "• 规范化问句统计（本课高频）\n• 依据不足问题列表\n• 提问留存窗口\n• 教师视图不展示学生身份\n• 反馈是「待核对线索」，不是成绩画像",
         TEAL_LIGHT, TEAL, TEAL, INK, 19, 16)
add_card(slide, 6.85, 1.80, 5.75, 4.15,
         "后续争取 / 诚实边界",
         "• 小规模同义问题组（默认关闭）\n• 阈值未校准，效果待真实条件验证\n• 不据此宣称教学质量提升\n• 与学生问答共用课程边界与留存策略\n• 演示时可打开教师侧反馈入口一眼带过",
         WHITE, LINE, AMBER, INK, 19, 16)

# 9. Gaps (three) + later plan teaser
slide = add_slide("现有不足（三条）", "把边界讲清楚，比夸大进展更有说服力", 9, total=TOTAL)
add_card(slide, 0.55, 1.80, 3.95, 4.25,
         "① 评价未人工复核",
         "仅有弱规则命中与拒答\n事实正确 ≠ 引用支持\n尚无独立人工事实正确率\n题集为 AI 预标",
         RED_LIGHT, RED, RED, INK, 18, 15)
add_card(slide, 4.70, 1.80, 3.95, 4.25,
         "② 复杂资料未覆盖",
         "非 OCR / 扫描件主线\n公式与图表理解有限\n纯图页覆盖不全\n文字层 PDF 是中期范围",
         AMBER_LIGHT, AMBER, AMBER, INK, 18, 15)
add_card(slide, 8.85, 1.80, 3.85, 4.25,
         "③ 产品化待验",
         "阈值与同义分组未校准\n供应商账单 / 成本待验\n单机演示，非校园网并发\n开题争取项尚未全部兑现",
         WHITE, LINE, NAVY, INK, 18, 15)

# 10. Later plan (four) + conclusion
slide = add_slide("后期计划与中期结论", "四条可汇报的下一步；不承诺公网生产部署", 10, total=TOTAL)
todo = [
    ("人工复核", "小样本抽检事实正确与引用支持是否一致"),
    ("检索对照", "Hybrid / rerank 与 dense 对照，有无收益都如实报"),
    ("同义问题组", "教师侧校准后再开，服务备课线索"),
    ("试用与成本", "结题非开发者试用 + 资源账单；不承诺公网"),
]
for i, (title, body) in enumerate(todo):
    x = 0.55 + (i % 2) * 6.35
    y = 1.75 + (i // 2) * 1.70
    add_card(slide, x, y, 6.05, 1.50, f"0{i+1} {title}", body, AMBER_LIGHT if i % 2 else TEAL_LIGHT, AMBER if i % 2 else TEAL, AMBER if i % 2 else TEAL, INK, 17, 14)
add_box(slide, 0.55, 5.35, 12.2, 1.05, NAVY, NAVY)
add_text(slide, "结论：双叙事亮点可讲清——课界+证据闭环，与页级检索/生成技术；本机可演示。开题「重点争取」项诚实后置。", 0.80, 5.55, 11.7, 0.65, 16, WHITE, True, PP_ALIGN.CENTER)

prs.core_properties.title = "高校课程知识库智能问答：中期进展（参考稿）"
prs.core_properties.subject = "中文四课 cn_v2、双叙事亮点、不足与后期计划"
prs.core_properties.author = "项目组"
prs.core_properties.keywords = "RAG, 课程知识库, 引用, 中期, 参考稿"
prs.save(PPTX_PATH)

speaker_notes = """# 中期汇报讲稿提示（参考稿 v5）

> **定位**：仓内 PPT / 本讲稿均为**参考稿**，队友可改；终稿可由答辩同学按《答辩提纲》重做。
> 建议时长：约 6–8 分钟。只讲已验证事实；不要把计划说成已完成。

1. **封面**：高校课程知识库问答；核心是学生能回到原始材料核对。标明参考稿。
2. **问题与目标**：课件分散难定位；通用模型流畅≠课程依据。中期交付本机文本型原型 + 四课 4/27/2501。
3. **业务闭环**：资料管理 → 授权检索 → 证据问答 → 来源核对 → 教学反馈。
4. **技术亮点（双叙事）**：先讲教学闭环（权限贯穿、可拒答、回原页、三角色+反馈）；再讲检索/生成（页级门禁、BGE-M3、deepseek-flash、策略可开关；主演示 dense、rerank 关）。
5. **技术栈一图**：本地实现课程边界与来源映射；开源复用，不自造平台。
6. **四课证据**：cn_v2；弱规则 12/12 与 4/4 ≠ 人工事实正确率。
7. **演示路径**：LA-01 有引用并打开 `#page=`；LA-R 依据不足；PERM 无权限可口头或实走。
8. **教学反馈**：高频问句与依据不足是线索，不画像学生能力。
9. **不足三条**：无人工复核；复杂 PDF/OCR 未覆盖；阈值/同义/账单与并发待验。
10. **后期四条**：人工复核；Hybrid/rerank 对照；同义问题组；结题试用+成本。不承诺公网。开题「重点争取」诚实后置。

演示前：库 `univ_course_rag_cn_v2`，`UPLOAD_DIR`→uploads_v2，`RERANKER_BACKEND=disabled`，模型 `deepseek-flash`。队友自备 Key。不要临场换模型或重新入库。
"""
NOTES_PATH.write_text(speaker_notes, encoding="utf-8")
HANDOFF_NOTES.parent.mkdir(parents=True, exist_ok=True)
HANDOFF_NOTES.write_text(speaker_notes, encoding="utf-8")
shutil.copy2(PPTX_PATH, HANDOFF_PPT)

check = Presentation(PPTX_PATH)
all_text = "\n".join(shape.text for slide in check.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))
required = [
    "参考稿", "2501", "cn_v2", "deepseek-flash", "Dense", "教学反馈",
    "线性代数", "业务闭环", "依据不足", "事实正确", "PyMuPDF", "BGE-M3",
    "双叙事", "教学场景闭环", "后期", "人工复核",
]
stale = [
    "Docling", "P7-3", "PHIL", "V4 Flash", "v4-flash", "r7b6",
    "已停止的路线", "已停止路线", "Frege", "心理学、线性代数、历史、哲学",
    "univ_course_rag_r9_restore", "RERANKER 关闭（残留",
]
missing = [item for item in required if item not in all_text]
stale_found = [item for item in stale if item in all_text]
bounds_errors = []
for si, slide in enumerate(check.slides, 1):
    for shape in slide.shapes:
        if shape.left < 0 or shape.top < 0 or shape.left + shape.width > check.slide_width + 1000 or shape.top + shape.height > check.slide_height + 1000:
            bounds_errors.append({"slide": si, "shape": shape.name})
images = sum(1 for slide in check.slides for shape in slide.shapes if shape.shape_type == 13)
handoff_same = HANDOFF_PPT.exists() and HANDOFF_PPT.read_bytes() == PPTX_PATH.read_bytes()
validation = {
    "created_at": datetime.now().astimezone().isoformat(),
    "file": str(PPTX_PATH.relative_to(ROOT)),
    "handoff_file": str(HANDOFF_PPT.relative_to(ROOT)),
    "handoff_bytes_match": handoff_same,
    "size_bytes": PPTX_PATH.stat().st_size,
    "slide_count": len(check.slides),
    "embedded_images": images,
    "required_claims_missing": missing,
    "stale_claims_found": stale_found,
    "geometry_bounds_errors": bounds_errors,
    "status": "passed" if len(check.slides) == TOTAL and not missing and not stale_found and not bounds_errors and handoff_same else "failed",
}
VALIDATION_PATH.write_text(json.dumps(validation, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(validation, ensure_ascii=False, indent=2))
if validation["status"] != "passed":
    raise SystemExit(1)
