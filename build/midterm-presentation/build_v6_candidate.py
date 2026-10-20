from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path.cwd()
BASE = ROOT / "build/midterm-presentation"
OUT = BASE / "output"
FINALIZER = BASE / "finalizer"
ASSETS = BASE / "assets-v6"
OUT.mkdir(parents=True, exist_ok=True)
FINALIZER.mkdir(parents=True, exist_ok=True)
ASSETS.mkdir(parents=True, exist_ok=True)

PPTX_PATH = OUT / "final-v6-candidate.pptx"
VALIDATION_PATH = FINALIZER / "final-v6-candidate.validation.json"
NOTES_PATH = OUT / "speaker-notes-v6-candidate.md"
HANDOFF_PPT = ROOT / "中期答辩材料/汇报PPT/中期汇报-v6-candidate.pptx"
HANDOFF_NOTES = ROOT / "中期答辩材料/汇报PPT/讲稿-v6-candidate.md"
HANDOFF_SHOTS = ROOT / "中期答辩材料/彩排截图"
DIAGRAMS = ROOT / "中期答辩材料/图示"

# Keep v5 paths untouched — never write to these:
V5_PPT = ROOT / "中期答辩材料/汇报PPT/中期汇报-v5.pptx"
V5_NOTES = ROOT / "中期答辩材料/汇报PPT/讲稿-v5.md"

W, H = 13.333, 7.5
FONT = "Microsoft YaHei"
NAVY = "163247"
TEAL = "247D78"
TEAL_LIGHT = "E4F2F0"
AMBER = "D39A37"
AMBER_LIGHT = "FBF0D8"
RED = "B84242"
RED_LIGHT = "F9E7E5"
GREEN = "26734E"
GREEN_LIGHT = "E4F1E9"
INK = "263642"
MUTED = "61727D"
LINE = "D6E0E5"
BG = "F5F8FA"
WHITE = "FFFFFF"

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]
TOTAL = 9


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return shape


def add_box(slide, x, y, w, h, fill=WHITE, stroke=LINE, radius=True, line_width=1.1):
    geometry = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(geometry, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(stroke)
    shape.line.width = Pt(line_width)
    return shape


def add_card(slide, x, y, w, h, title, body, fill=WHITE, stroke=LINE, title_color=NAVY, body_color=INK, title_size=19, body_size=15):
    shape = add_box(slide, x, y, w, h, fill, stroke)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.16)
    tf.margin_top = Inches(0.10)
    tf.margin_bottom = Inches(0.08)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(6)
    r = p.add_run()
    r.text = title
    r.font.name = FONT
    r.font.size = Pt(title_size)
    r.font.bold = True
    r.font.color.rgb = rgb(title_color)
    for line in body.split("\n"):
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(2)
        r = p.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(body_size)
        r.font.color.rgb = rgb(body_color)
    return shape


def add_label(slide, text, x, y, w, fill=TEAL_LIGHT, color=TEAL):
    add_box(slide, x, y, w, 0.38, fill, fill, radius=True, line_width=0.5)
    add_text(slide, text, x + 0.05, y + 0.015, w - 0.10, 0.34, 13, color, True, PP_ALIGN.CENTER)


def add_line(slide, x1, y1, x2, y2, color=TEAL, width=2.0):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    return shape


def add_header(slide, title, subtitle, page, total=TOTAL):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)
    add_text(slide, title, 0.62, 0.28, 12.0, 0.48, 26, NAVY, True)
    add_text(slide, subtitle, 0.66, 0.82, 11.9, 0.36, 13, MUTED)
    add_box(slide, 0.66, 1.28, 0.72, 0.045, AMBER, AMBER, radius=False, line_width=0.2)
    add_text(slide, "高校课程知识库智能问答 · 中期汇报候选稿 v6", 0.62, 7.12, 9.8, 0.20, 9, MUTED)
    add_text(slide, f"{page:02d} / {total:02d}", 11.7, 7.10, 0.95, 0.22, 9, MUTED, align=PP_ALIGN.RIGHT)


def add_slide(title, subtitle, page):
    slide = prs.slides.add_slide(blank)
    add_header(slide, title, subtitle, page)
    return slide


def add_stat(slide, x, y, w, value, label, color=TEAL, h=1.05):
    add_box(slide, x, y, w, h, WHITE, LINE)
    add_text(slide, value, x + 0.06, y + 0.08, w - 0.12, 0.48, 22, color, True, PP_ALIGN.CENTER)
    add_text(slide, label, x + 0.06, y + 0.58, w - 0.12, 0.36, 12, MUTED, align=PP_ALIGN.CENTER)


def add_image_contain(slide, path: Path, x, y, w, h):
    with Image.open(path) as im:
        ratio = im.width / im.height
    out_w = min(w, h * ratio)
    out_h = out_w / ratio
    left = x + (w - out_w) / 2
    top = y + (h - out_h) / 2
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(out_w), height=Inches(out_h))


def focus_shot(src: Path, dest: Path, max_h: int = 1000) -> Path:
    if not src.exists():
        raise FileNotFoundError(f"Missing screenshot: {src}")
    with Image.open(src) as original:
        top = 0
        bottom = original.height
        if original.height > max_h + 100:
            top = min(40, original.height - max_h)
            bottom = min(original.height, top + max_h)
        original.crop((0, top, original.width, bottom)).convert("RGB").save(dest, quality=90)
    return dest


# ---------------------------------------------------------------------------
# 1. Cover
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = rgb(NAVY)
add_box(slide, 0, 0, 6.7, H, NAVY, NAVY, radius=False, line_width=0.1)
add_box(slide, 7.05, 1.45, 5.55, 4.55, "1D4053", "416173", radius=True, line_width=1.0)
add_card(slide, 7.35, 1.80, 2.35, 1.20, "课界权限", "未授权拒绝\n截图 10 · PERM", TEAL_LIGHT, TEAL, TEAL, INK, 16, 12)
add_card(slide, 10.00, 1.80, 2.30, 1.20, "证据链", "引用回原页\n#page=", WHITE, WHITE, NAVY, INK, 16, 12)
add_line(slide, 9.75, 2.40, 9.95, 2.40, AMBER, 2.0)
add_card(slide, 7.35, 3.30, 2.35, 1.20, "失败语义", "有据 / 不足\n故障 / 权限", AMBER_LIGHT, AMBER, AMBER, INK, 16, 12)
add_card(slide, 10.00, 3.30, 2.30, 1.20, "可审计评测", "22 题分层\n开发者核验", WHITE, WHITE, NAVY, INK, 16, 12)
add_card(slide, 8.20, 4.75, 3.20, 0.95, "演示库 cn_v2", "4 课 · 27 PDF · 2501 chunks", GREEN_LIGHT, GREEN, GREEN, INK, 15, 12)
add_text(slide, "有依据、可核对、能解释失败", 7.30, 5.95, 5.0, 0.32, 15, "D8E7EE", True, PP_ALIGN.CENTER)
add_text(slide, "高校课程知识库\n智能问答", 0.78, 1.70, 5.6, 1.45, 36, WHITE, True)
add_box(slide, 0.82, 3.40, 0.9, 0.06, AMBER, AMBER, radius=False, line_width=0.1)
add_text(slide, "中期汇报 · 候选稿 v6", 0.82, 3.70, 5.2, 0.48, 24, "EAF2F6", True)
add_text(slide, "有依据、可核对、能解释失败", 0.84, 4.35, 5.3, 0.36, 16, "D3E1E8")
add_text(slide, "候选稿：保留 v5 原稿；终稿可由答辩同学改写", 0.84, 5.70, 5.5, 0.36, 13, AMBER, True)
add_text(slide, "中文四课文本型 RAG · Dense 主演示 · 2026.10", 0.84, 6.40, 5.5, 0.34, 13, "D3E1E8")

# ---------------------------------------------------------------------------
# 2. Pain → goal
# ---------------------------------------------------------------------------
slide = add_slide("痛点 → 中期目标", "课件难定位；流畅 ≠ 课程依据；需要可核对且能解释失败的问答", 2)
add_card(
    slide, 0.60, 1.55, 5.95, 4.55,
    "实际痛点",
    "• 知识分散在多份讲义/课件中\n• 自然语言问题难靠文件名定位\n• 通用模型流畅，未必对应当前课程口径\n• 学生需要原文核对入口\n• 多课程必须守住访问边界\n• 「答得漂亮」不等于「课内有据」",
    WHITE, LINE, NAVY, INK, 18, 15,
)
add_card(
    slide, 6.80, 1.55, 5.85, 4.55,
    "中期目标（本阶段）",
    "• 本机可演示的文本型 RAG 原型\n• 中文四课：4 / 27 / 2,501（cn_v2）\n• 学生：提问 → 引用 → 回原 PDF 页\n• 依据不足 / 系统失败 / 权限拒绝分开\n• 教师侧基础反馈（线索，非画像）\n• 不宣称生产部署或通用多模态",
    AMBER_LIGHT, AMBER, AMBER, INK, 18, 15,
)
add_box(slide, 0.60, 6.25, 12.05, 0.55, NAVY, NAVY)
add_text(
    slide,
    "核心问题：如何让课程问答遵守课界、提供可核对依据，并在依据不足时正确停止？",
    0.80, 6.32, 11.7, 0.42, 15, WHITE, True, PP_ALIGN.CENTER,
)

# ---------------------------------------------------------------------------
# 3. Closed loop
# ---------------------------------------------------------------------------
slide = add_slide("方案闭环", "资料 → 授权检索 → 证据问答 → 来源核对 → 教学反馈；三角色分工", 3)
steps = [
    (0.50, "① 资料", "上传解析入库\n页级质量态", TEAL_LIGHT, TEAL),
    (3.00, "② 授权检索", "课程成员边界\n课内候选召回", WHITE, TEAL),
    (5.50, "③ 证据问答", "受限上下文\n有据/拒答/失败", AMBER_LIGHT, AMBER),
    (8.00, "④ 来源核对", "引用摘录\n打开原 PDF 页", WHITE, AMBER),
    (10.50, "⑤ 教学反馈", "高频问句\n依据不足线索", GREEN_LIGHT, GREEN),
]
for i, (x, title, body, fill, stroke) in enumerate(steps):
    add_card(slide, x, 1.55, 2.25, 2.40, title, body, fill, stroke, stroke if fill != WHITE else NAVY, INK, 16, 13)
    if i < len(steps) - 1:
        add_line(slide, x + 2.28, 2.70, steps[i + 1][0] - 0.08, 2.70, AMBER, 2.0)

arch = DIAGRAMS / "架构总览.png"
add_box(slide, 0.55, 4.20, 7.40, 2.45, WHITE, LINE)
if arch.exists():
    add_image_contain(slide, arch, 0.70, 4.30, 7.10, 2.25)
else:
    add_text(slide, "（架构图缺失）", 0.80, 5.10, 6.8, 0.40, 14, MUTED, align=PP_ALIGN.CENTER)

add_card(
    slide, 8.15, 4.20, 4.50, 2.45,
    "三角色",
    "管理员：账号 / 课程 / 授权\n教师：资料状态 + 基础反馈\n学生：提问 → 引用 → 回原页\n\n不是单一聊天窗",
    TEAL_LIGHT, TEAL, TEAL, INK, 16, 13,
)

# ---------------------------------------------------------------------------
# 4. Highlight 1 — course boundary
# ---------------------------------------------------------------------------
slide = add_slide("亮点① 课程边界贯穿", "未授权请求被拒绝；权限不是 UI 装饰（依据：PERM 截图 + 运行卡）", 4)
perm = HANDOFF_SHOTS / "10-PERM-无权限.png"
perm_img = ASSETS / "perm-focus.png"
if perm.exists():
    focus_shot(perm, perm_img)
    add_box(slide, 0.55, 1.50, 6.90, 4.85, WHITE, LINE)
    add_image_contain(slide, perm_img, 0.70, 1.65, 6.60, 4.40)
    add_text(slide, "彩排截图 10 · PERM 无权限", 0.70, 6.10, 6.60, 0.22, 12, MUTED, align=PP_ALIGN.CENTER)
else:
    add_card(slide, 0.55, 1.50, 6.90, 4.85, "PERM 截图", "缺少 10-PERM-无权限.png", RED_LIGHT, RED, RED, INK, 18, 14)

add_card(
    slide, 7.70, 1.50, 4.95, 4.85,
    "机制与边界",
    "解决什么\n未加入课程仍可问答会泄漏资料\n\n关键机制\n课程成员校验贯穿检索/\n生成/打开原文件\n未授权 → 拒绝（403/无权限 UI）\n\n现有证据\n截图 10；运行卡 PERM\n历史 API/UI 验收\n\n边界（勿夸大）\n仅覆盖已测路径；\n不宣称已过渗透/并发审计",
    WHITE, LINE, TEAL, INK, 17, 13,
)

# ---------------------------------------------------------------------------
# 5. Highlight 2 — evidence chain
# ---------------------------------------------------------------------------
slide = add_slide("亮点② 证据链到原文", "页级解析 → 课内检索 → 受限生成与引用 → 点开 #page= 原 PDF", 5)
shot03 = ASSETS / "ask-cite.png"
shot04 = ASSETS / "open-pdf.png"
focus_shot(HANDOFF_SHOTS / "03-学生-线代作答与引用.png", shot03)
focus_shot(HANDOFF_SHOTS / "04-学生-引用打开PDF.png", shot04)

add_box(slide, 0.50, 1.45, 4.55, 3.55, WHITE, LINE)
add_image_contain(slide, shot03, 0.60, 1.55, 4.35, 3.05)
add_text(slide, "03 回答与引用", 0.60, 4.70, 4.35, 0.22, 12, NAVY, True, PP_ALIGN.CENTER)

add_box(slide, 5.20, 1.45, 4.55, 3.55, WHITE, LINE)
add_image_contain(slide, shot04, 5.30, 1.55, 4.35, 3.05)
add_text(slide, "04 真浏览器打开 PDF", 5.30, 4.70, 4.35, 0.22, 12, NAVY, True, PP_ALIGN.CENTER)

chain = DIAGRAMS / "问答链路.png"
add_box(slide, 9.90, 1.45, 2.85, 3.55, WHITE, LINE)
if chain.exists():
    add_image_contain(slide, chain, 9.98, 1.55, 2.70, 3.15)
else:
    add_text(slide, "问答链路图", 10.0, 2.90, 2.6, 0.40, 12, MUTED, align=PP_ALIGN.CENTER)

add_label(slide, "页级解析/质量态", 0.50, 5.15, 2.95, TEAL_LIGHT, TEAL)
add_label(slide, "Dense · BGE-M3", 3.60, 5.15, 2.70, TEAL_LIGHT, TEAL)
add_label(slide, "deepseek-flash", 6.45, 5.15, 2.70, AMBER_LIGHT, AMBER)
add_label(slide, "引用 #page=", 9.30, 5.15, 3.40, GREEN_LIGHT, GREEN)
add_text(
    slide,
    "只引用实际送入上下文的来源；检索命中 ≠ 答案正确。主演示 Dense，rerank 关。",
    0.55, 5.75, 12.2, 0.35, 14, MUTED, align=PP_ALIGN.CENTER,
)

# ---------------------------------------------------------------------------
# 6. Highlight 3 — failure semantics
# ---------------------------------------------------------------------------
slide = add_slide("亮点③ 失败语义对照", "有据 ≠ 依据不足 ≠ 系统故障 ≠ 权限拒绝；无证据不等于坏了，也不等于答案正确", 6)
cells = [
    (0.50, "有据回答", "课内有证据\n附引用可回页\n例：LA-01", GREEN_LIGHT, GREEN),
    (3.70, "依据不足", "无足够课内依据\n正确停止/拒答\n例：LA-R · 截图 05", AMBER_LIGHT, AMBER),
    (6.90, "系统故障", "上游/服务异常\n可重试，非拒答\n截图 11", RED_LIGHT, RED),
    (10.10, "权限拒绝", "未授权课程\n403 / 无权限 UI\n截图 10", TEAL_LIGHT, TEAL),
]
for x, title, body, fill, stroke in cells:
    add_card(slide, x, 1.50, 2.95, 2.55, title, body, fill, stroke, stroke, INK, 17, 13)

shot05 = ASSETS / "refuse.png"
shot11 = ASSETS / "sysfail.png"
focus_shot(HANDOFF_SHOTS / "05-学生-拒答.png", shot05)
focus_shot(HANDOFF_SHOTS / "11-系统故障-重试.png", shot11)
add_box(slide, 0.50, 4.25, 6.00, 2.35, WHITE, LINE)
add_image_contain(slide, shot05, 0.60, 4.35, 5.80, 1.95)
add_text(slide, "05 依据不足", 0.60, 6.35, 5.80, 0.20, 11, AMBER, True, PP_ALIGN.CENTER)
add_box(slide, 6.70, 4.25, 6.00, 2.35, WHITE, LINE)
add_image_contain(slide, shot11, 6.80, 4.35, 5.80, 1.95)
add_text(slide, "11 系统故障（可重试）", 6.80, 6.35, 5.80, 0.20, 11, RED, True, PP_ALIGN.CENTER)

# ---------------------------------------------------------------------------
# 7. Highlight 4 — audit / gates
# ---------------------------------------------------------------------------
slide = add_slide("亮点④ 可审计评测与门禁", "分层指标 · verifier=developer · 22 题不可外推 · 不合成单一准确率", 7)
add_stat(slide, 0.50, 1.45, 2.40, "4/27/2501", "课 / PDF / chunks", TEAL)
add_stat(slide, 3.05, 1.45, 2.55, "20/27", "COVERAGE 需核对\n≠ 解析失败", AMBER)
add_stat(slide, 5.75, 1.45, 2.40, "Dense", "主演示 · rerank关", NAVY)
add_stat(slide, 8.30, 1.45, 2.25, "BGE-M3", "Embedding", TEAL)
add_stat(slide, 10.70, 1.45, 2.05, "flash", "deepseek-flash", AMBER)

# metrics table as cards
add_card(
    slide, 0.50, 2.70, 8.00, 3.55,
    "22 题分层结果（Dense / Hybrid）",
    "端到端可答题　　15/17　/　16/17\n已答事实要点　　15/16　/　16/16\n已答引用支持　　16/16　/　16/16\n严格预期页命中　12/16　/　12/16　← 不叫引用错误率\n明显课外拒答　　5/5　　/　5/5　　← 不可泛化困难拒答\n\n来源：研究证据/results/retrieval-compare-2026-10-20.md\n不得宣称 Hybrid 显著更优；Keyword-only→最终 evidence=0",
    WHITE, LINE, NAVY, INK, 16, 13,
)
add_card(
    slide, 8.70, 2.70, 4.05, 3.55,
    "真实限制",
    "PROB-F02\n应答题却 NO_EVIDENCE\n根因未定位\n\n页码口径\n阅读器页序 ≠\nChunk 合并末页\n\n开发者核验\n非独立评测",
    RED_LIGHT, RED, RED, INK, 15, 12,
)

# ---------------------------------------------------------------------------
# 8. Demo path 60–90s
# ---------------------------------------------------------------------------
slide = add_slide("60–90 秒主演示路径", "每步证明什么：有据 → 回原页 → 依据不足 →（截图）权限 / 系统失败", 8)
demo_steps = [
    ("1 提问 LA-01", "课内可检索\n并生成有据答", TEAL_LIGHT, TEAL),
    ("2 看引用", "来源可核对\n摘录+页码", WHITE, TEAL),
    ("3 开 PDF", "#page= 回原页\n截图 03+04", AMBER_LIGHT, AMBER),
    ("4 LA-R 拒答", "依据不足停止\n截图 05", WHITE, AMBER),
    ("5 权限/故障", "PERM≠故障\n截图 10+11", GREEN_LIGHT, GREEN),
]
for i, (title, body, fill, stroke) in enumerate(demo_steps):
    x = 0.45 + i * 2.55
    add_card(slide, x, 1.50, 2.40, 2.35, title, body, fill, stroke, stroke if fill != WHITE else NAVY, INK, 15, 12)
    if i < len(demo_steps) - 1:
        add_line(slide, x + 2.42, 2.55, x + 2.52, 2.55, AMBER, 1.8)

add_box(slide, 0.50, 4.10, 12.25, 2.45, WHITE, LINE)
add_text(slide, "口播要点（固定问句）", 0.75, 4.25, 11.5, 0.32, 16, TEAL, True)
add_text(
    slide,
    "LA-01：根据线性代数讲义，线性代数主要研究哪两类对象？ → 有回答与引用\n"
    "LA-R：根据线性代数讲义，如何用 SQL 创建数据库表并插入数据？ → 依据不足，不编造\n"
    "PERM：未加入课程 → 无权限（截图 10）｜系统故障：可重试、文案不同于依据不足（截图 11）\n"
    "运行卡：docs/中期演示运行卡.md　｜　截图：中期答辩材料/彩排截图/",
    0.75, 4.70, 11.7, 1.55, 14, INK,
)

# ---------------------------------------------------------------------------
# 9. Limits + later + conclusion
# ---------------------------------------------------------------------------
slide = add_slide("局限 · 后期 · 中期结论", "把边界讲清；不承诺公网生产；开题原件未找到则诚实说明", 9)
add_card(
    slide, 0.50, 1.45, 4.00, 3.55,
    "现有不足",
    "• 开发者核验，非独立评测\n• 页码口径仍需核对\n• 无 OCR / 扫描主线\n• 公式与图表理解有限\n• 阈值未校准；成本未对账\n• PROB-F02 根因未定位",
    RED_LIGHT, RED, RED, INK, 16, 13,
)
add_card(
    slide, 4.70, 1.45, 4.00, 3.55,
    "后期方向",
    "• 独立小样本人工复核\n• 单题诊断 PROB-F02\n• Hybrid/rerank 有无收益如实报\n• 同义问题组（校准后再开）\n• 结题试用 + 资源账单\n• 不承诺公网部署",
    AMBER_LIGHT, AMBER, AMBER, INK, 16, 13,
)
add_card(
    slide, 8.90, 1.45, 3.85, 3.55,
    "开题对照",
    "• 获批开题原件：未找到\n• 仅有修订稿可对\n• 扫描/公式/图：\n  未声称已完整实现\n• 「重点争取」项后置",
    WHITE, LINE, NAVY, INK, 16, 13,
)
add_box(slide, 0.50, 5.20, 12.25, 1.35, NAVY, NAVY)
add_text(
    slide,
    "中期结论：四课本机可演示——课界贯穿、证据可回原页、失败语义可区分、评测可分层审计。\n"
    "检索命中 ≠ 答案正确；22 题不可外推；候选稿 v6，终稿可由答辩同学改写。",
    0.75, 5.40, 11.8, 1.00, 15, WHITE, True, PP_ALIGN.CENTER,
)

prs.core_properties.title = "高校课程知识库智能问答：中期汇报候选稿 v6"
prs.core_properties.subject = "技术亮点、分层评测、失败语义、演示路径"
prs.core_properties.author = "项目组"
prs.core_properties.keywords = "RAG, 课程知识库, 引用, 中期, 候选稿 v6"
prs.save(PPTX_PATH)

speaker_notes = """# 中期汇报讲稿（候选稿 v6）

> **定位**：本讲稿与 `中期汇报-v6-candidate.pptx` 配套，为**候选稿**；`讲稿-v5.md` / `中期汇报-v5.pptx` **保留不动**。终稿可由答辩同学改写。
> **建议时长**：约 6–8 分钟。只讲已验证事实；不把计划说成已完成；不合成单一准确率。

证据索引：`中期答辩材料/彩排截图/`（03/04/05/10/11）、`中期答辩材料/图示/`、`研究证据/results/retrieval-compare-2026-10-20.md`、`docs/中期演示运行卡.md`。

---

## 第 1 页 · 封面（约 20s）

高校课程知识库问答。副题：**有依据、可核对、能解释失败**。标明候选稿 v6。四格预告：课界权限、证据链、失败语义、可审计评测。库口径 cn_v2：4/27/2501。

## 第 2 页 · 痛点 → 目标（约 45s）

痛点：课件分散难定位；通用模型流畅 ≠ 课程依据；多课必须守边界。  
**核心叙事问题**（见 `研究证据/07`）：如何让课程问答遵守课界、提供可核对依据，并在依据不足时正确停止？  
中期交付：本机文本型原型；学生提问→引用→回原页；三种失败状态分开；不宣称生产/多模态。

## 第 3 页 · 方案闭环（约 40s）

五步：资料 → 授权检索 → 证据问答 → 来源核对 → 教学反馈。三角色：管理员/教师/学生。可指架构总览图一眼带过，少念栈名。

## 第 4 页 · 亮点① 课程边界（约 50s）

痛点：未授权仍问答会泄漏。机制：成员校验贯穿检索/生成/打开原文件。证据：截图 **10-PERM-无权限.png** + 运行卡 PERM。边界：只谈已测路径，不夸渗透审计。

## 第 5 页 · 亮点② 证据链到原文（约 55s）

链路：页级解析与质量态 → Dense（BGE-M3）→ deepseek-flash 受限生成与引用 → `#page=`。证据：截图 **03**、**04**（真浏览器开 PDF）；可选问答链路图。口播：**检索命中 ≠ 答案正确**；主演示 Dense，rerank 关。

## 第 6 页 · 亮点③ 失败语义对照（约 55s）

四格：有据 / 依据不足 / 系统故障 / 权限拒绝。用 LA-01 口播 + 截图 **05**（依据不足）、**11**（系统故障可重试）、**10**（权限）。强调：**无证据 ≠ 坏了 ≠ 答案正确**。

## 第 7 页 · 亮点④ 可审计评测与门禁（约 70s）

规模 4/27/2501；COVERAGE_INCOMPLETE=部分页需核对（**≠** 解析失败）。分层指标（**原样念，勿合成准确率**）：

| 指标 | Dense | Hybrid |
| --- | --- | --- |
| 端到端可答题 | 15/17 | 16/17 |
| 已答事实 | 15/16 | 16/16 |
| 已答引用支持 | 16/16 | 16/16 |
| 严格预期页 | 12/16 | 12/16（不叫引用错误率） |
| 明显课外拒答 | 5/5 | 5/5 |

verifier=developer；22 题不可外推；不得宣称 Hybrid 显著更优；Keyword-only→最终 evidence=0。**PROB-F02**：应答题却 NO_EVIDENCE，根因未定位——作为真实限制讲。

## 第 8 页 · 60–90s 主演示路径（约 60–90s 实操或口播+截图）

| 步 | 动作 | 证明什么 |
| --- | --- | --- |
| 1 | LA-01 提问 | 课内可检索并有据生成 |
| 2 | 看引用区 | 来源可核对 |
| 3 | 点开 PDF `#page=` | 证据链到原文（03+04） |
| 4 | LA-R | 依据不足正确停止（05） |
| 5 | 截图/口播 PERM 与系统故障 | 权限拒绝 ≠ 系统故障（10+11） |

固定问句见运行卡。无实机时用彩排截图演练，不宣称已录屏（若仓内尚无录屏文件）。

## 第 9 页 · 局限 + 后期 + 结论（约 50s）

不足：开发者核验；页码口径；无 OCR；公式/图有限；阈值/成本待验；PROB-F02。  
后期：独立复核；单题诊断；Hybrid/rerank 如实对照；同义组；试用+账单；**不承诺公网**。  
开题原件未找到则诚实一句。  
结论：四课本机可演示，四亮点可讲清；候选稿，终稿可改。

---

演示前：库 `univ_course_rag_cn_v2`，`UPLOAD_DIR`→uploads_v2，`RERANKER_BACKEND=disabled`，模型 `deepseek-flash`。队友自备 Key。不要临场换模型或重新入库。
"""
NOTES_PATH.write_text(speaker_notes, encoding="utf-8")
HANDOFF_NOTES.parent.mkdir(parents=True, exist_ok=True)
HANDOFF_NOTES.write_text(speaker_notes, encoding="utf-8")
shutil.copy2(PPTX_PATH, HANDOFF_PPT)

# Guard: v5 must still exist and we must not have overwritten them via this script's outputs
assert HANDOFF_PPT.resolve() != V5_PPT.resolve()
assert HANDOFF_NOTES.resolve() != V5_NOTES.resolve()
if not V5_PPT.exists() or not V5_NOTES.exists():
    raise SystemExit("v5 handoff files missing — abort")

check = Presentation(PPTX_PATH)
all_text = "\n".join(
    shape.text for slide in check.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False)
)
required = [
    "候选稿", "2501", "cn_v2", "deepseek-flash", "Dense", "15/17", "16/17",
    "12/16", "5/5", "PROB-F02", "依据不足", "系统故障", "PERM",
    "BGE-M3", "COVERAGE", "Keyword-only", "#page=",
]
stale = [
    "Docling", "PHIL", "V4 Flash", "v4-flash", "Frege",
    "univ_course_rag_r9_restore", "准确率百分之", "显著优于 Dense",
    "真实模型核验尚未运行",
]
missing = [item for item in required if item not in all_text]
stale_found = [item for item in stale if item in all_text]
bounds_errors = []
for si, slide in enumerate(check.slides, 1):
    for shape in slide.shapes:
        if (
            shape.left < 0
            or shape.top < 0
            or shape.left + shape.width > check.slide_width + 1000
            or shape.top + shape.height > check.slide_height + 1000
        ):
            bounds_errors.append({"slide": si, "shape": shape.name})
images = sum(1 for slide in check.slides for shape in slide.shapes if shape.shape_type == 13)
handoff_same = HANDOFF_PPT.exists() and HANDOFF_PPT.read_bytes() == PPTX_PATH.read_bytes()
v5_bytes = V5_PPT.stat().st_size
v5_notes_ok = V5_NOTES.exists()
validation = {
    "created_at": datetime.now().astimezone().isoformat(),
    "file": str(PPTX_PATH.relative_to(ROOT)),
    "handoff_file": str(HANDOFF_PPT.relative_to(ROOT)),
    "handoff_bytes_match": handoff_same,
    "v5_pptx_still_present": V5_PPT.exists(),
    "v5_pptx_size_bytes": v5_bytes,
    "v5_notes_still_present": v5_notes_ok,
    "size_bytes": PPTX_PATH.stat().st_size,
    "slide_count": len(check.slides),
    "embedded_images": images,
    "required_claims_missing": missing,
    "stale_claims_found": stale_found,
    "geometry_bounds_errors": bounds_errors,
    "status": (
        "passed"
        if (
            len(check.slides) == TOTAL
            and 8 <= len(check.slides) <= 10
            and not missing
            and not stale_found
            and not bounds_errors
            and handoff_same
            and V5_PPT.exists()
            and v5_notes_ok
        )
        else "failed"
    ),
}
VALIDATION_PATH.write_text(json.dumps(validation, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(validation, ensure_ascii=False, indent=2))
if validation["status"] != "passed":
    raise SystemExit(1)
